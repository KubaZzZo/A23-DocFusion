"""Campus procurement demo for DocFusion.

Runs a deterministic end-to-end flow without requiring the desktop UI:
parse texts, extract entities, fuse cross-document entities, fill a template,
and format a generated DOCX with a natural-language command.
"""

from __future__ import annotations

import asyncio
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from core.doc_commander import DocCommander
from core.entity_extractor import EntityExtractor
from core.template_filler import TemplateFiller
from core.workflow_engine import CrossDocFusionStep, WorkflowContext
from db.database import DocumentDAO, EntityDAO
from db.models import configure_database, init_db, reset_database


SAMPLES = [
    {
        "filename": "contract_a.txt",
        "text": (
            "校园采购合同A\n"
            "甲方：金陵科技学院创新中心\n"
            "乙方：北京智远科技有限公司\n"
            "合同金额：268,000元\n"
            "签订日期：2024-05-12\n"
            "联系人：张伟\n"
            "联系电话：13800138000\n"
            "邮箱：zhangwei@example.com\n"
            "地址：北京市海淀区中关村大街1号\n"
            "统一社会信用代码：913100001234567890\n"
        ),
    },
    {
        "filename": "contract_b.txt",
        "text": (
            "校园采购合同B\n"
            "采购方：金陵科技学院创新中心\n"
            "供应商：北京智远科技\n"
            "预算金额：310,000元\n"
            "签约日期：2024-06-03\n"
            "负责人：李明\n"
            "联系方式：13900139000\n"
            "联系邮箱：liming@example.com\n"
            "通讯地址：北京市海淀区中关村大街1号\n"
            "项目编号：913100001234567890\n"
        ),
    },
    {
        "filename": "invoice_c.txt",
        "text": (
            "采购发票记录C\n"
            "购买方：金陵科技学院创新中心\n"
            "销售方：北京智远科技有限公司\n"
            "价税合计：268,000元\n"
            "开票日期：2024-05-20\n"
            "经办人：王芳\n"
            "电话：13700137000\n"
            "电子邮箱：wangfang@example.com\n"
            "供应商地址：北京市海淀区中关村大街1号\n"
            "纳税人识别号：913100001234567890\n"
        ),
    },
]


class DemoLLM:
    """Small deterministic LLM replacement for demo entity extraction."""

    async def extract_json(self, prompt: str, user_input: str) -> dict:
        entities = EntityExtractor._extract_regex_entities(user_input)
        known_values = [
            ("organization", "金陵科技学院创新中心"),
            ("organization", "北京智远科技有限公司"),
            ("organization", "北京智远科技"),
            ("person", "张伟"),
            ("person", "李明"),
            ("person", "王芳"),
            ("address", "北京市海淀区中关村大街1号"),
        ]
        for entity_type, value in known_values:
            if value in user_input:
                entities.append(
                    {
                        "type": entity_type,
                        "value": value,
                        "context": value,
                        "confidence": 0.95,
                    }
                )
        return {"entities": entities, "summary": "校园采购资料", "topic": "校园采购"}


async def main() -> None:
    with tempfile.TemporaryDirectory(prefix="docfusion_demo_") as temp_dir:
        temp = Path(temp_dir)
        configure_database(f"sqlite:///{temp / 'demo.db'}")
        init_db()
        try:
            await run_demo(temp)
        finally:
            reset_database()


async def run_demo(temp: Path) -> None:
    print("=" * 60)
    print("DocFusion 校园采购合同审查与汇总演示")
    print("=" * 60)

    extractor = EntityExtractor(llm_client=DemoLLM(), enable_verify=False)
    doc_ids: list[int] = []

    print("\n[1/5] 导入并抽取实体")
    for sample in SAMPLES:
        source = temp / sample["filename"]
        source.write_text(sample["text"], encoding="utf-8")
        doc = DocumentDAO.create(source.name, "txt", str(source))
        DocumentDAO.update_text(doc.id, sample["text"])
        result = await extractor.extract(sample["text"], force=True)
        EntityDAO.create_batch(doc.id, result["entities"])
        doc_ids.append(doc.id)
        print(f"  - {source.name}: {len(result['entities'])} entities")

    print("\n[2/5] 跨文档融合")
    ctx = WorkflowContext()
    fusion = await CrossDocFusionStep(doc_ids).run(ctx)
    cross_doc = fusion["cross_doc"]
    for item in cross_doc[:5]:
        variants = ", ".join(item.get("variants", []))
        print(f"  - {item['type']} {item['value']} in {item['doc_count']} docs; variants: {variants}")

    entities = []
    for doc_id in doc_ids:
        for entity in EntityDAO.get_by_document(doc_id):
            entities.append({"type": entity.entity_type, "value": entity.entity_value, "confidence": entity.confidence})

    print("\n[3/5] 自动填写汇总表")
    template_path = create_template(temp)
    fill_result = await TemplateFiller(llm_client=DemoLLM()).fill(str(template_path), entities)
    print(f"  - accuracy: {fill_result['accuracy']:.0%}")
    print(f"  - output: {fill_result['output_path']}")

    print("\n[4/5] 自然语言格式化 DOCX")
    docx_path = create_docx_summary(temp)
    command_result = await DocCommander(use_codex=False, llm_client=DemoLLM()).execute_command(
        "把第一段加粗并居中显示",
        str(docx_path),
    )
    print(f"  - command success: {command_result.get('success')}")
    print(f"  - formatted docx: {docx_path}")

    print("\n[5/5] 演示完成")
    print(f"  - fused groups: {len(cross_doc)}")
    print(f"  - filled table: {fill_result['output_path']}")
    print(f"  - formatted document: {docx_path}")


def create_template(temp: Path) -> Path:
    path = temp / "procurement_summary_template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "采购汇总"
    headers = ["甲方", "供应商", "合同金额", "签订日期", "联系人", "电话", "邮箱", "地址", "税号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=1, column=index, value=header)
        ws.cell(row=2, column=index, value="")
    wb.save(path)
    return path


def create_docx_summary(temp: Path) -> Path:
    path = temp / "procurement_summary.docx"
    doc = Document()
    doc.add_paragraph("校园采购数据融合汇总")
    doc.add_paragraph("本报告由 DocFusion 演示脚本生成。")
    doc.save(path)
    return path


if __name__ == "__main__":
    asyncio.run(main())
