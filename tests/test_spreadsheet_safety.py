"""Spreadsheet formula injection hardening tests."""
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.spreadsheet_safety import escape_formula_value
from core.template_filler import TemplateFiller

TMP_DIR = Path(__file__).parent / ".tmp_spreadsheet_safety"


def test_escape_formula_value_prefixes_dangerous_strings():
    assert escape_formula_value("=cmd|calc") == "'=cmd|calc"
    assert escape_formula_value("+SUM(A1:A2)") == "'+SUM(A1:A2)"
    assert escape_formula_value("-10+20") == "'-10+20"
    assert escape_formula_value("@HYPERLINK") == "'@HYPERLINK"
    assert escape_formula_value("safe") == "safe"
    assert escape_formula_value(42) == 42


def test_template_filler_escapes_xlsx_formula_values():
    shutil.rmtree(TMP_DIR, ignore_errors=True)
    TMP_DIR.mkdir(exist_ok=True)
    path = TMP_DIR / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "name"
    workbook.save(path)

    TemplateFiller()._fill_xlsx(
        str(path),
        [{"field_name": "name", "sheet": "Sheet", "row": 2, "col": 1}],
        {"name": "=cmd|calc"},
    )

    loaded = load_workbook(path)
    assert loaded.active["A2"].value == "'=cmd|calc"
    loaded.close()
    shutil.rmtree(TMP_DIR, ignore_errors=True)
