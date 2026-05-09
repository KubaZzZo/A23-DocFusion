from pathlib import Path
from uuid import uuid4

from docx import Document

from core.template_filler import TemplateFiller


TEST_DATA_DIR = Path(__file__).parent / "test_data"


def test_fill_docx_writes_values_into_table_cells():
    path = TEST_DATA_DIR / f"fill_docx_{uuid4().hex}.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Phone"
    doc.save(path)

    try:
        analysis = TemplateFiller()._analyze_docx(path)
        TemplateFiller()._fill_docx(str(path), analysis["fields"], {"Name": "Alice", "Phone": "010"})
        updated = Document(path)

        assert updated.tables[0].cell(1, 0).text == "Alice"
        assert updated.tables[0].cell(1, 1).text == "010"
    finally:
        path.unlink(missing_ok=True)
