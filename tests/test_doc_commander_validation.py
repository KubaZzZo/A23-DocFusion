"""DocCommander command validation tests."""
from pathlib import Path
from uuid import uuid4

from docx import Document

from core.doc_commander import BACKUP_DIR, DocCommander

TEST_DATA_DIR = Path(__file__).parent / "test_data"


def _make_docx(name: str = "validation") -> Path:
    path = TEST_DATA_DIR / f"{name}_working_{uuid4().hex}.docx"
    doc = Document()
    doc.add_paragraph("原始内容")
    doc.save(path)
    return path


def _backup_names() -> set[str]:
    BACKUP_DIR.mkdir(exist_ok=True)
    return {p.name for p in BACKUP_DIR.glob("*.docx")}


def test_execute_rejects_invalid_format_color_before_backup():
    path = _make_docx()
    before_backups = _backup_names()
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(
            str(path),
            {"action": "format", "target": "paragraph", "params": {"index": 0, "color": [255, 0]}},
        )

        assert result["success"] is False
        assert "color" in result["message"]
        assert _backup_names() == before_backups
    finally:
        path.unlink(missing_ok=True)


def test_execute_rejects_invalid_edit_operation_before_backup():
    path = _make_docx()
    before_backups = _backup_names()
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(
            str(path),
            {"action": "edit", "params": {"operation": "rewrite_all", "text": "新内容"}},
        )

        assert result["success"] is False
        assert "operation" in result["message"]
        assert _backup_names() == before_backups
    finally:
        path.unlink(missing_ok=True)


def test_execute_rejects_empty_find_replace_before_backup():
    path = _make_docx()
    before_backups = _backup_names()
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(str(path), {"action": "find_replace", "params": {"find": "", "replace": "x"}})

        assert result["success"] is False
        assert "find" in result["message"]
        assert _backup_names() == before_backups
    finally:
        path.unlink(missing_ok=True)


def test_find_replace_updates_table_cells():
    path = TEST_DATA_DIR / f"table_replace_{uuid4().hex}.docx"
    doc = Document()
    doc.add_paragraph("outside old")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "inside old"
    doc.save(path)
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(
            str(path),
            {"action": "find_replace", "params": {"find": "old", "replace": "new"}},
        )

        updated = Document(path)
        assert result["success"] is True
        assert updated.paragraphs[0].text == "outside new"
        assert updated.tables[0].cell(0, 0).text == "inside new"
    finally:
        path.unlink(missing_ok=True)


def test_execute_rejects_invalid_extract_target_without_backup():
    path = _make_docx()
    before_backups = _backup_names()
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(str(path), {"action": "extract", "target": "comments", "params": {}})

        assert result["success"] is False
        assert "target" in result["message"]
        assert _backup_names() == before_backups
    finally:
        path.unlink(missing_ok=True)


def test_execute_rejects_missing_action_before_backup():
    path = _make_docx()
    before_backups = _backup_names()
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(str(path), {"target": "paragraph", "params": {"index": 0, "bold": True}})

        assert result["success"] is False
        assert "action" in result["message"]
        assert _backup_names() == before_backups
    finally:
        path.unlink(missing_ok=True)


def test_execute_rejects_unknown_action_before_backup():
    path = _make_docx()
    before_backups = _backup_names()
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(str(path), {"action": "rewrite", "params": {}})

        assert result["success"] is False
        assert "action" in result["message"]
        assert _backup_names() == before_backups
    finally:
        path.unlink(missing_ok=True)


def test_execute_rejects_unexpected_command_field_before_backup():
    path = _make_docx()
    before_backups = _backup_names()
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(
            str(path),
            {"action": "format", "target": "paragraph", "params": {"index": 0}, "unsafe": True},
        )

        assert result["success"] is False
        assert "unsafe" in result["message"]
        assert _backup_names() == before_backups
    finally:
        path.unlink(missing_ok=True)


def test_execute_rejects_unexpected_format_param_before_backup():
    path = _make_docx()
    before_backups = _backup_names()
    commander = DocCommander.__new__(DocCommander)

    try:
        result = commander.execute(
            str(path),
            {"action": "format", "target": "paragraph", "params": {"index": 0, "macro": "run"}},
        )

        assert result["success"] is False
        assert "macro" in result["message"]
        assert _backup_names() == before_backups
    finally:
        path.unlink(missing_ok=True)


def test_execute_edit_extract_and_structure_handlers():
    path = TEST_DATA_DIR / f"handler_case_{uuid4().hex}.docx"
    doc = Document()
    doc.add_paragraph("old paragraph")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Header"
    table.cell(1, 0).text = "Value"
    doc.save(path)
    commander = DocCommander.__new__(DocCommander)

    try:
        edit_result = commander.execute(
            str(path),
            {"action": "edit", "params": {"operation": "replace", "index": 0, "text": "new paragraph"}},
        )
        structure_result = commander.execute(
            str(path),
            {"action": "structure", "params": {"operation": "add_heading", "text": "Added Heading", "level": 2}},
        )
        extract_text = commander.execute(str(path), {"action": "extract", "target": "text", "params": {}})
        extract_tables = commander.execute(str(path), {"action": "extract", "target": "tables", "params": {}})
        updated = Document(path)

        assert edit_result["success"] is True
        assert structure_result["success"] is True
        assert updated.paragraphs[0].text == "new paragraph"
        assert any(p.text == "Added Heading" for p in updated.paragraphs)
        assert "new paragraph" in extract_text["data"]
        assert extract_tables["data"] == [[["Header"], ["Value"]]]
    finally:
        path.unlink(missing_ok=True)
