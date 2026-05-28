import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.doc_commander import DocCommander
from core.entity_extractor import EntityExtractor
from core.semantic_matcher import SemanticMatcher
from core.template_filler import TemplateFiller


class FakeExtractLLM:
    model = "gpt-4o"

    async def extract_json(self, prompt, text):
        return {
            "entities": [
                {"type": "organization", "value": "Acme Corp", "context": text[:40], "confidence": 0.95},
                {"type": "amount", "value": "1200000", "context": text[:40], "confidence": 0.92},
            ],
            "summary": "contract summary",
            "topic": "contract",
        }


@pytest.mark.asyncio
async def test_entity_extractor_with_mock_llm_merges_normalized_entities():
    extractor = EntityExtractor(enable_verify=False)
    extractor.llm = FakeExtractLLM()

    result = await extractor.extract("Acme Corp signed contract amount 1200000. Contact ok@example.com", force=True)

    values = {entity["value"] for entity in result["entities"]}
    assert "Acme Corp" in values
    assert "1200000" in values
    assert "ok@example.com" in values
    assert result["summary"] == "contract summary"


@pytest.mark.asyncio
async def test_semantic_matcher_uses_local_rules_before_llm():
    class FailingLLM:
        async def extract_json(self, prompt, text):
            raise AssertionError("local matches should not call LLM")

    matcher = SemanticMatcher()
    matcher.llm = FailingLLM()

    result = await matcher.match(
        ["email", "e-mail"],
        [{"type": "email", "value": "contact@example.com", "confidence": 0.96}],
    )

    assert [match["field"] for match in result["matches"]] == ["email", "e-mail"]
    assert result["unmatched_fields"] == []


@pytest.mark.asyncio
async def test_template_filler_analyzes_and_fills_xlsx_with_mock_matcher(tmp_path, monkeypatch):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["email", "amount"])
    sheet["A2"] = ""
    sheet["B2"] = ""
    workbook.save(template)

    class FakeMatcher:
        async def match(self, fields, entities):
            return {
                "matches": [
                    {"field": "email", "value": "contact@example.com"},
                    {"field": "amount", "value": "=1+1"},
                ],
                "unmatched_fields": [],
            }

    output_dir = tmp_path / "outputs"
    output_dir.mkdir()
    monkeypatch.setattr("core.template_filler.OUTPUT_DIR", output_dir)
    filler = TemplateFiller()
    filler.matcher = FakeMatcher()

    result = await filler.fill(str(template), [])

    assert result["success"] is True
    assert result["filled"] == 2
    filled = load_workbook(result["output_path"])
    try:
        row = next(filled.active.iter_rows(min_row=2, max_row=2, values_only=True))
        assert row == ("contact@example.com", "'=1+1")
    finally:
        filled.close()


@pytest.mark.asyncio
async def test_doc_commander_rule_based_format_command_executes_on_docx(tmp_path):
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("first line")
    document.save(path)

    commander = DocCommander()
    command = await commander.parse_command("make first paragraph bold and font size 20", "docx")
    result = commander.execute(str(path), command)

    assert result["success"] is True
    changed = Document(str(path))
    run = changed.paragraphs[0].runs[0]
    assert run.bold is True
    assert run.font.size.pt == 20
