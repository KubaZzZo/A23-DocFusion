import asyncio
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from db.models import configure_database, reset_database, init_db
from core.document_workflow import DocumentWorkflow


@pytest.fixture()
def isolated_backend(tmp_path):
    configure_database(f"sqlite:///{tmp_path / 'docfusion.db'}")
    init_db()
    try:
        yield tmp_path
    finally:
        reset_database()


def _docx_bytes(path: Path) -> bytes:
    doc = Document()
    doc.add_paragraph("first line")
    doc.add_paragraph("second line")
    doc.save(path)
    return path.read_bytes()


def test_document_command_creates_downloadable_version_and_can_rollback(isolated_backend):
    workflow = DocumentWorkflow(upload_dir=isolated_backend / "uploads")
    uploaded = workflow.upload_document("versioned.docx", _docx_bytes(isolated_backend / "source.docx"))
    workflow.parse_document(uploaded["id"])

    result = asyncio.run(workflow.execute_command(uploaded["id"], "make the first paragraph bold"))

    assert result["result"]["success"] is True
    versions = workflow.list_versions(uploaded["id"])
    assert len(versions) == 1
    assert versions[0]["version_no"] == 1
    assert Path(versions[0]["file_path"]).is_file()

    changed = Document(uploaded["path"])
    assert changed.paragraphs[0].runs[0].bold is True

    rollback = workflow.rollback_version(uploaded["id"], versions[0]["id"])
    assert rollback["success"] is True
    restored = Document(uploaded["path"])
    assert restored.paragraphs[0].runs[0].bold is None
