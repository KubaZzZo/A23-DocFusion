import sys
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from api.server import app
from db.database import DocumentDAO, DocumentVersionDAO, EntityDAO, TemplateDAO
from db.models import configure_database, init_db, reset_database


def _headers():
    return {"Authorization": "Bearer competition-test-token"}


def setup_function():
    pass


def teardown_function():
    reset_database()


def _configure(tmp_path, monkeypatch):
    configure_database(f"sqlite:///{tmp_path / 'docfusion.db'}")
    monkeypatch.setenv("DOCFUSION_API_TOKEN", "competition-test-token")
    init_db()


def test_entity_graph_api_returns_document_entity_network(tmp_path, monkeypatch):
    _configure(tmp_path, monkeypatch)
    first = DocumentDAO.create("contract-a.docx", "docx", str(tmp_path / "a.docx"))
    second = DocumentDAO.create("contract-b.docx", "docx", str(tmp_path / "b.docx"))
    EntityDAO.create_batch(first.id, [{"type": "organization", "value": "Acme", "confidence": 0.95}])
    EntityDAO.create_batch(second.id, [{"type": "organization", "value": "Acme", "confidence": 0.8}])

    response = TestClient(app).get("/api/entities/graph", headers=_headers())

    assert response.status_code == 200
    payload = response.json()
    assert any(node["kind"] == "entity" and node["label"] == "Acme" for node in payload["nodes"])
    assert len(payload["edges"]) == 2


def test_template_review_and_confirmed_fill(tmp_path, monkeypatch):
    _configure(tmp_path, monkeypatch)
    doc = DocumentDAO.create("contract.docx", "docx", str(tmp_path / "contract.docx"))
    EntityDAO.create_batch(
        doc.id,
        [
            {"type": "organization", "value": "Acme Corp", "confidence": 0.92},
            {"type": "amount", "value": "1200000", "confidence": 0.9},
        ],
    )
    template_path = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["organization", "amount"])
    sheet.append([None, None])
    workbook.save(template_path)
    template = TemplateDAO.create(
        "template.xlsx",
        str(template_path),
        '{"field_names":["organization","amount"],"fields":[{"field_name":"organization","sheet":"Sheet","row":2,"col":1},{"field_name":"amount","sheet":"Sheet","row":2,"col":2}]}',
    )

    client = TestClient(app)
    review = client.post(
        "/api/templates/fill/review",
        json={"template_id": template.id, "document_ids": [doc.id]},
        headers=_headers(),
    )
    assert review.status_code == 200
    suggestions = {item["field"]: item["suggested_value"] for item in review.json()["suggestions"]}
    assert suggestions == {"organization": "Acme Corp", "amount": "1200000"}

    confirmed = client.post(
        "/api/templates/fill/confirmed",
        json={"template_id": template.id, "fill_map": suggestions},
        headers=_headers(),
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["filled"] == 2
    assert confirmed.json()["result_download_url"].startswith("/api/outputs/")


def test_template_review_marks_low_confidence_matches_for_manual_review(tmp_path, monkeypatch):
    _configure(tmp_path, monkeypatch)
    doc = DocumentDAO.create("contract.docx", "docx", str(tmp_path / "contract.docx"))
    EntityDAO.create_batch(
        doc.id,
        [{"type": "organization", "value": "Possible Corp", "confidence": 0.62}],
    )
    template_path = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["organization"])
    sheet.append([None])
    workbook.save(template_path)
    template = TemplateDAO.create(
        "template.xlsx",
        str(template_path),
        '{"field_names":["organization"],"fields":[{"field_name":"organization","sheet":"Sheet","row":2,"col":1}]}',
    )

    review = TestClient(app).post(
        "/api/templates/fill/review",
        json={"template_id": template.id, "document_ids": [doc.id]},
        headers=_headers(),
    )

    assert review.status_code == 200
    suggestion = review.json()["suggestions"][0]
    assert suggestion["suggested_value"] == "Possible Corp"
    assert suggestion["review_required"] is True
    assert suggestion["status"] == "needs_review"


def test_demo_report_and_document_diff_apis(tmp_path, monkeypatch):
    _configure(tmp_path, monkeypatch)
    client = TestClient(app)

    demo = client.post("/api/demo/load", headers=_headers())
    assert demo.status_code == 200
    assert demo.json()["documents"] >= 3

    report = client.get("/api/reports/full", headers=_headers())
    assert report.status_code == 200
    assert report.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert len(report.content) > 1000

    pdf_report = client.get("/api/reports/full?format=pdf", headers=_headers())
    assert pdf_report.status_code == 200
    assert pdf_report.headers["content-type"].startswith("application/pdf")
    assert pdf_report.content.startswith(b"%PDF")
    assert len(pdf_report.content) > 500

    doc_path = tmp_path / "diff.docx"
    document = Document()
    document.add_paragraph("before text")
    document.save(doc_path)
    doc = DocumentDAO.create("diff.docx", "docx", str(doc_path))
    old_path = tmp_path / "before.docx"
    Document(str(doc_path)).save(old_path)
    changed = Document(str(doc_path))
    changed.paragraphs[0].text = "after text"
    changed.save(doc_path)
    version = DocumentVersionDAO.create(doc.id, str(old_path), "replace before with after")

    diff = client.get(f"/api/documents/{doc.id}/versions/{version.id}/diff", headers=_headers())
    assert diff.status_code == 200
    payload = diff.json()
    assert "before text" in payload["before"]
    assert "after text" in payload["after"]
    assert "-before text" in payload["diff"]
