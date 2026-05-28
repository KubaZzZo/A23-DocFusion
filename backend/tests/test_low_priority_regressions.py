import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.doc_commander import DocCommander
from core.file_signature import validate_file_signature
from core.text_chunker import TextChunker
from core.workflow_errors import WorkflowValidationError
from crawler.news_spider import NewsSpider
from db.database import DocumentDAO, EntityDAO
from db.models import configure_database, init_db, reset_database


def test_text_chunker_prefers_paragraph_boundaries_for_overlap():
    text = "alpha paragraph\nbeta paragraph\ncharlie paragraph"

    chunks = TextChunker.chunk(text, chunk_size=30, overlap=5)

    assert chunks == ["alpha paragraph\nbeta paragraph", "charlie paragraph"]


def test_edit_replace_preserves_paragraph_style(tmp_path):
    path = tmp_path / "styled.docx"
    doc = Document()
    paragraph = doc.add_paragraph("old value")
    paragraph.style = "Heading 1"
    paragraph.alignment = 1
    doc.save(path)

    result = DocCommander()._handle_edit(
        str(path),
        {"operation": "replace", "index": 0, "text": "new value"},
    )

    assert result["success"] is True
    changed = Document(path)
    assert changed.paragraphs[0].text == "new value"
    assert changed.paragraphs[0].style.name == "Heading 1"
    assert changed.paragraphs[0].alignment == 1


def test_cross_document_entities_preserve_filenames_with_commas(tmp_path):
    configure_database(f"sqlite:///{tmp_path / 'docfusion.db'}")
    init_db()
    try:
        first = DocumentDAO.create("alpha,inc.docx", "docx", str(tmp_path / "a.docx"))
        second = DocumentDAO.create("beta.docx", "docx", str(tmp_path / "b.docx"))
        EntityDAO.create_batch(first.id, [{"type": "organization", "value": "Acme", "context": "", "confidence": 0.9}])
        EntityDAO.create_batch(second.id, [{"type": "organization", "value": "Acme", "context": "", "confidence": 0.8}])

        rows = EntityDAO.get_cross_document_entities(min_documents=2)

        assert rows[0]["documents"] == ["alpha,inc.docx", "beta.docx"]
    finally:
        reset_database()


@pytest.mark.parametrize("filename", ["notes.txt", "notes.md"])
def test_text_upload_signature_rejects_binary_payload(filename):
    with pytest.raises(WorkflowValidationError):
        validate_file_signature(filename, b"\x00\x01\x02\x03not text")


def test_news_spider_follows_redirects():
    spider = NewsSpider()
    try:
        assert spider.client.follow_redirects is True
    finally:
        spider.client.close()


def test_gitignore_excludes_nested_database_files():
    ignore = Path(__file__).resolve().parents[2] / ".gitignore"
    patterns = ignore.read_text(encoding="utf-8")

    assert "**/*.db" in patterns
    assert "**/*.sqlite" in patterns
