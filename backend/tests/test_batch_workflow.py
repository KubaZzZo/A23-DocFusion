import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from db.models import configure_database, reset_database, init_db
from core.batch_workflow import BatchWorkflow


@pytest.fixture()
def isolated_backend(tmp_path):
    configure_database(f"sqlite:///{tmp_path / 'docfusion.db'}")
    init_db()
    try:
        yield tmp_path
    finally:
        reset_database()


def test_batch_workflow_uploads_parses_extracts_and_exports_report(isolated_backend):
    def fake_extract(text: str) -> dict:
        return {
            "entities": [
                {"type": "email", "value": "alpha@example.com", "context": text, "confidence": 0.95},
                {"type": "phone", "value": "13800138000", "context": text, "confidence": 0.95},
            ],
            "summary": "demo summary",
        }

    workflow = BatchWorkflow(
        upload_dir=isolated_backend / "uploads",
        output_dir=isolated_backend / "outputs",
        extractor=fake_extract,
    )

    result = workflow.process_files(
        [
            ("alpha.txt", "email alpha@example.com phone 13800138000".encode("utf-8")),
            ("beta.txt", "email beta@example.com".encode("utf-8")),
        ]
    )

    assert result["total"] == 2
    assert result["succeeded"] == 2
    assert result["failed"] == 0
    assert result["entities_count"] == 4
    assert Path(result["report_path"]).is_file()

    workbook = load_workbook(result["report_path"])
    assert workbook.sheetnames == ["文档汇总", "实体明细"]
    summary_rows = list(workbook["文档汇总"].iter_rows(values_only=True))
    entity_rows = list(workbook["实体明细"].iter_rows(values_only=True))
    assert summary_rows[0] == ("文档ID", "文件名", "类型", "解析字数", "实体数", "状态", "错误")
    assert len(summary_rows) == 3
    assert entity_rows[0] == ("文档ID", "文件名", "实体类型", "实体值", "置信度", "上下文")
    assert len(entity_rows) == 5
