"""模板表格填写器"""
import json
import shutil
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from docx import Document as DocxDocument
from core.semantic_matcher import SemanticMatcher
from core.spreadsheet_safety import escape_formula_value
from config import OUTPUT_DIR


class TemplateFiller:
    def __init__(self, provider: str = None, llm_client=None):
        self.matcher = SemanticMatcher(provider, llm_client=llm_client)

    async def analyze_template(self, file_path: str) -> dict:
        """分析模板，识别需要填写的字段"""
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return self._analyze_xlsx(path)
        elif suffix == ".docx":
            return self._analyze_docx(path)
        else:
            raise ValueError(f"不支持的模板格式: {suffix}")

    def _analyze_xlsx(self, path: Path) -> dict:
        """分析xlsx模板，支持横向（首行表头）和纵向（首列标签）布局"""
        wb = load_workbook(str(path))
        fields = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # 检测布局模式：横向（首行有多个非空）vs 纵向（首列有多个非空）
            first_row_count = sum(1 for cell in ws[1] if cell.value)
            first_col_count = sum(1 for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=1)
                                  for cell in row if cell.value)

            is_vertical = first_col_count > first_row_count

            if is_vertical:
                # 纵向布局：首列=标签，次列=值
                for row in ws.iter_rows(min_row=1, values_only=False):
                    label_cell = row[0] if len(row) > 0 else None
                    value_cell = row[1] if len(row) > 1 else None
                    if label_cell and label_cell.value and value_cell:
                        if value_cell.value is None or str(value_cell.value).strip() == "":
                            fields.append({
                                "field_name": str(label_cell.value).strip(),
                                "sheet": sheet_name,
                                "row": value_cell.row,
                                "col": value_cell.column,
                            })
            else:
                # 横向布局：首行=表头
                headers = []
                for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=False), []):
                    if cell.value:
                        headers.append({"col": cell.column, "name": str(cell.value), "sheet": sheet_name})

                for row in ws.iter_rows(min_row=2, values_only=False):
                    for cell in row:
                        if cell.value is None or str(cell.value).strip() == "":
                            header = next((h for h in headers if h["col"] == cell.column), None)
                            if header:
                                fields.append({
                                    "field_name": header["name"],
                                    "sheet": sheet_name,
                                    "row": cell.row,
                                    "col": cell.column,
                                })
        wb.close()
        return {"fields": fields, "field_names": list(set(f["field_name"] for f in fields))}

    def _analyze_docx(self, path: Path) -> dict:
        """分析docx模板：表格（横向/纵向）+ 正文占位符{{field}}"""
        doc = DocxDocument(str(path))
        fields = []

        # 1. 扫描表格
        for t_idx, table in enumerate(doc.tables):
            if not table.rows:
                continue

            # 检测布局：横向 vs 纵向
            first_row_filled = sum(1 for cell in table.rows[0].cells if cell.text.strip())
            first_col_filled = sum(1 for row in table.rows if row.cells and row.cells[0].text.strip())
            is_vertical = first_col_filled > first_row_filled

            if is_vertical:
                # 纵向：首列=标签，次列=值
                for r_idx, row in enumerate(table.rows):
                    if len(row.cells) < 2:
                        continue
                    label = row.cells[0].text.strip()
                    value = row.cells[1].text.strip()
                    if label and not value:
                        fields.append({
                            "field_name": label,
                            "table_index": t_idx,
                            "row": r_idx,
                            "col": 1,
                        })
            else:
                # 横向：首行=表头
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                for r_idx, row in enumerate(table.rows[1:], start=1):
                    for c_idx, cell in enumerate(row.cells):
                        if not cell.text.strip() and c_idx < len(headers) and headers[c_idx]:
                            fields.append({
                                "field_name": headers[c_idx],
                                "table_index": t_idx,
                                "row": r_idx,
                                "col": c_idx,
                            })

        # 2. 扫描正文占位符 {{field_name}}
        import re
        placeholder_pattern = re.compile(r"\{\{([^}]+)\}\}")
        for p_idx, paragraph in enumerate(doc.paragraphs):
            for match in placeholder_pattern.finditer(paragraph.text):
                field_name = match.group(1).strip()
                fields.append({
                    "field_name": field_name,
                    "paragraph_index": p_idx,
                    "placeholder": match.group(0),
                })

        return {"fields": fields, "field_names": list(set(f["field_name"] for f in fields))}

    async def fill(self, template_path: str, entities: list[dict]) -> dict:
        """自动填写模板"""
        # 1. 分析模板
        analysis = await self.analyze_template(template_path)
        field_names = analysis["field_names"]

        if not field_names:
            return {"success": False, "message": "模板中未找到需要填写的字段"}

        # 2. 语义匹配
        match_result = await self.matcher.match(field_names, entities)
        matches = match_result.get("matches", [])

        # 3. 构建填写映射
        fill_map = {}
        for m in matches:
            fill_map[m["field"]] = m["value"]

        # 4. 执行填写
        path = Path(template_path)
        suffix = path.suffix.lower()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_name = f"{path.stem}_filled_{timestamp}{suffix}"
        output_path = OUTPUT_DIR / output_name
        shutil.copyfile(template_path, output_path)
        # 移除不安全的chmod调用

        if suffix == ".xlsx":
            self._fill_xlsx(str(output_path), analysis["fields"], fill_map)
        elif suffix == ".docx":
            self._fill_docx(str(output_path), analysis["fields"], fill_map)

        field_names = analysis["field_names"]
        filled_count = sum(1 for field_name in field_names if field_name in fill_map)
        total_count = len(field_names)
        filled_cells = sum(1 for f in analysis["fields"] if f["field_name"] in fill_map)
        total_cells = len(analysis["fields"])

        return {
            "success": True,
            "output_path": str(output_path),
            "filled": filled_count,
            "total": total_count,
            "accuracy": filled_count / total_count if total_count > 0 else 0,
            "filled_cells": filled_cells,
            "total_cells": total_cells,
            "message": f"{filled_count}/{total_count} fields matched, written to {filled_cells} cells",
            "unmatched": match_result.get("unmatched_fields", []),
        }

    def _fill_xlsx(self, file_path: str, fields: list[dict], fill_map: dict):
        wb = load_workbook(file_path)
        for f in fields:
            value = fill_map.get(f["field_name"])
            if value:
                ws = wb[f["sheet"]]
                ws.cell(row=f["row"], column=f["col"], value=escape_formula_value(value))
        wb.save(file_path)
        wb.close()

    def _fill_docx(self, file_path: str, fields: list[dict], fill_map: dict):
        doc = DocxDocument(file_path)
        for f in fields:
            value = fill_map.get(f["field_name"])
            if value:
                if "table_index" in f:
                    # 表格单元格填充
                    table = doc.tables[f["table_index"]]
                    table.rows[f["row"]].cells[f["col"]].text = value
                elif "paragraph_index" in f:
                    # 正文占位符替换
                    para = doc.paragraphs[f["paragraph_index"]]
                    para.text = para.text.replace(f["placeholder"], value)
        doc.save(file_path)
