"""Document parser supporting docx/md/xlsx/txt/pdf/image formats."""
import subprocess
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Callable

from config import OCR_CONFIG
from logger import get_logger

log = get_logger("core.parser")


class ParserAdapter:
    """Small adapter wrapper for one document format family."""

    def __init__(self, suffixes: set[str], parse_func: Callable[[Path], str]):
        self.suffixes = suffixes
        self.parse_func = parse_func

    def parse(self, path: Path) -> str:
        return self.parse_func(path)


class DocumentParser:
    IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".bmp"}
    SUPPORTED_TYPES = {".docx", ".doc", ".md", ".xlsx", ".txt", ".pdf"} | IMAGE_TYPES
    _ADAPTERS: dict[str, ParserAdapter] = {}

    @staticmethod
    def parse(file_path: str) -> dict:
        """Parse a document and return text, file type, and metadata."""
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix not in DocumentParser.SUPPORTED_TYPES:
            raise ValueError(f"Unsupported file format: {suffix}")

        stat = path.stat() if path.exists() else None
        file_size = stat.st_size if stat else 0
        mtime = stat.st_mtime if stat else 0.0

        cache_before = DocumentParser._do_parse.cache_info()
        parsed = DocumentParser._do_parse(str(path.resolve()), mtime, file_size)
        cache_after = DocumentParser._do_parse.cache_info()

        return {
            "text": parsed["text"],
            "file_type": parsed["file_type"],
            "metadata": {
                **parsed["metadata"],
                "parsed_at": datetime.now().isoformat(),
                "cache_hit": cache_after.hits > cache_before.hits,
            },
        }

    @staticmethod
    @lru_cache(maxsize=64)
    def _do_parse(resolved_path: str, mtime: float, file_size: int) -> dict:
        path = Path(resolved_path)
        suffix = path.suffix.lower()
        metadata = {
            "filename": path.name,
            "file_size": file_size,
        }

        try:
            adapter = DocumentParser._get_adapter(suffix)
            text = adapter.parse(path)
        except Exception as e:
            log.error(f"Failed to parse file {path.name}: {e}")
            raise

        log.info(f"Parsed file: {path.name}, {len(text)} chars")
        return {"text": text, "file_type": suffix.lstrip("."), "metadata": metadata}

    @staticmethod
    def clear_cache():
        DocumentParser._do_parse.cache_clear()

    @classmethod
    def register_adapter(cls, adapter: ParserAdapter):
        for suffix in adapter.suffixes:
            cls._ADAPTERS[suffix] = adapter
        cls.SUPPORTED_TYPES.update(adapter.suffixes)

    @classmethod
    def _get_adapter(cls, suffix: str) -> ParserAdapter:
        if not cls._ADAPTERS:
            cls._register_default_adapters()
        adapter = cls._ADAPTERS.get(suffix)
        if not adapter:
            raise ValueError(f"Unsupported file format: {suffix}")
        return adapter

    @classmethod
    def _register_default_adapters(cls):
        cls.register_adapter(ParserAdapter({".txt"}, lambda path: cls._parse_txt(path)))
        cls.register_adapter(ParserAdapter({".md"}, lambda path: cls._parse_md(path)))
        cls.register_adapter(ParserAdapter({".docx"}, lambda path: cls._parse_docx(path)))
        cls.register_adapter(ParserAdapter({".doc"}, lambda path: cls._parse_doc(path)))
        cls.register_adapter(ParserAdapter({".xlsx"}, lambda path: cls._parse_xlsx(path)))
        cls.register_adapter(ParserAdapter({".pdf"}, lambda path: cls._parse_pdf(path)))
        cls.register_adapter(ParserAdapter(cls.IMAGE_TYPES, lambda path: cls._parse_image(path)))

    @staticmethod
    def _parse_md(path: Path) -> str:
        return path.read_text(encoding="utf-8")

    @staticmethod
    def _parse_txt(path: Path) -> str:
        return path.read_text(encoding="utf-8")

    @staticmethod
    def _parse_docx(path: Path) -> str:
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    @staticmethod
    def _parse_doc(path: Path) -> str:
        """Parse legacy .doc format using subprocess converters."""
        try:
            result = subprocess.run(
                ["antiword", str(path)],
                capture_output=True,
                text=True,
                timeout=30,
                check=False,
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout.strip()
        except (FileNotFoundError, subprocess.TimeoutExpired):
            pass

        try:
            result = subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "txt:Text", "--outdir", "/tmp", str(path)],
                capture_output=True,
                text=True,
                timeout=60,
                check=False,
            )
            if result.returncode == 0:
                txt_path = Path("/tmp") / f"{path.stem}.txt"
                if txt_path.exists():
                    content = txt_path.read_text(encoding="utf-8", errors="ignore")
                    txt_path.unlink()
                    return content.strip()
        except (FileNotFoundError, subprocess.TimeoutExpired):
            pass

        log.warning(".doc parsing requires antiword or libreoffice. Install: apt install antiword or libreoffice")
        return "[.doc parsing requires antiword or libreoffice]"

    @staticmethod
    def _parse_xlsx(path: Path) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"[Sheet {sheet_name}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        wb.close()
        return "\n".join(parts)

    @staticmethod
    def _parse_pdf(path: Path) -> str:
        try:
            import fitz  # PyMuPDF

            text_parts = []
            with fitz.open(str(path)) as doc:
                for page in doc:
                    page_text = page.get_text()
                    if page_text.strip():
                        text_parts.append(page_text)
            return "\n".join(text_parts)
        except ImportError:
            log.warning("PyMuPDF is not installed. Run: pip install PyMuPDF")
            return "[PDF parsing requires PyMuPDF]"

    @staticmethod
    def _parse_image(path: Path) -> str:
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            log.warning("pytesseract or Pillow is not installed. Run: pip install pytesseract Pillow")
            return "[Image OCR requires pytesseract and Pillow]"

        tesseract_cmd = OCR_CONFIG.get("tesseract_cmd")
        if tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

        try:
            with Image.open(path) as image:
                return pytesseract.image_to_string(image, lang=OCR_CONFIG.get("lang", "chi_sim+eng")).strip()
        except pytesseract.TesseractNotFoundError as e:
            msg = f"Tesseract-OCR not found, check OCR_CONFIG.tesseract_cmd: {e}"
            log.error(msg)
            raise RuntimeError(msg) from e
