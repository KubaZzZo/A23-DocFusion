"""Competition demo data and report workflows."""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook

from db.database import DocumentDAO, EntityDAO, TemplateDAO


class DemoWorkflow:
    def __init__(self, upload_dir: Path):
        self.upload_dir = Path(upload_dir)
        self.upload_dir.mkdir(exist_ok=True)

    def load_demo_data(self) -> dict:
        samples = [
            {
                "filename": "demo_contract_alpha.txt",
                "text": "甲方：星河科技\n乙方：云岭数据\n合同金额：120万元\n签署日期：2026-03-15\n项目：数据融合平台",
                "entities": [
                    ("organization", "星河科技"),
                    ("organization", "云岭数据"),
                    ("amount", "120万元"),
                    ("date", "2026-03-15"),
                    ("project", "数据融合平台"),
                ],
            },
            {
                "filename": "demo_contract_beta.txt",
                "text": "采购方：星河科技\n供应商：北辰智能\n合同金额：86万元\n签署日期：2026-04-02\n项目：智能填报系统",
                "entities": [
                    ("organization", "星河科技"),
                    ("organization", "北辰智能"),
                    ("amount", "86万元"),
                    ("date", "2026-04-02"),
                    ("project", "智能填报系统"),
                ],
            },
            {
                "filename": "demo_meeting_minutes.txt",
                "text": "评审会议确认星河科技与云岭数据共同推进数据融合平台，验收节点为2026-05-20。",
                "entities": [
                    ("organization", "星河科技"),
                    ("organization", "云岭数据"),
                    ("date", "2026-05-20"),
                    ("project", "数据融合平台"),
                ],
            },
        ]

        created_docs = 0
        created_entities = 0
        for sample in samples:
            path = self.upload_dir / sample["filename"]
            path.write_text(sample["text"], encoding="utf-8")
            doc = DocumentDAO.create(sample["filename"], "txt", str(path))
            DocumentDAO.update_text(doc.id, sample["text"])
            EntityDAO.create_batch(
                doc.id,
                [
                    {
                        "type": entity_type,
                        "value": value,
                        "context": sample["text"],
                        "confidence": 0.95,
                    }
                    for entity_type, value in sample["entities"]
                ],
            )
            created_docs += 1
            created_entities += len(sample["entities"])

        template_path = self.upload_dir / "demo_review_template.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "DemoTemplate"
        sheet.append(["organization", "amount", "date", "project"])
        sheet.append([None, None, None, None])
        workbook.save(template_path)
        TemplateDAO.create(
            "demo_review_template.xlsx",
            str(template_path),
            '{"field_names":["organization","amount","date","project"],"fields":[{"field_name":"organization","sheet":"DemoTemplate","row":2,"col":1},{"field_name":"amount","sheet":"DemoTemplate","row":2,"col":2},{"field_name":"date","sheet":"DemoTemplate","row":2,"col":3},{"field_name":"project","sheet":"DemoTemplate","row":2,"col":4}]}',
        )
        return {"documents": created_docs, "entities": created_entities, "templates": 1}


class ReportWorkflow:
    def build_full_report(self, fmt: str = "docx") -> bytes:
        if fmt.lower() == "pdf":
            return self._build_pdf_report()
        return self._build_docx_report()

    def _build_docx_report(self) -> bytes:
        documents = DocumentDAO.get_all()
        entities = EntityDAO.get_all()
        cross_rows = EntityDAO.get_cross_document_entities(min_documents=2, limit=100)
        type_counts = Counter(entity.entity_type for entity in entities)

        doc = DocxDocument()
        doc.add_heading("DocFusion 全流程演示报告", level=1)
        doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        doc.add_heading("文档列表", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "ID"
        table.rows[0].cells[1].text = "文件名"
        table.rows[0].cells[2].text = "类型"
        table.rows[0].cells[3].text = "已解析"
        for item in documents:
            row = table.add_row().cells
            row[0].text = str(item.id)
            row[1].text = item.filename
            row[2].text = item.file_type
            row[3].text = "是" if item.raw_text else "否"

        doc.add_heading("实体统计", level=2)
        for entity_type, count in sorted(type_counts.items()):
            doc.add_paragraph(f"{entity_type}: {count}", style=None)

        doc.add_heading("融合结果", level=2)
        if cross_rows:
            fusion = doc.add_table(rows=1, cols=5)
            headers = ["类型", "实体", "文档数", "次数", "关联文档"]
            for index, header in enumerate(headers):
                fusion.rows[0].cells[index].text = header
            for item in cross_rows:
                row = fusion.add_row().cells
                row[0].text = str(item.get("type", ""))
                row[1].text = str(item.get("value", ""))
                row[2].text = str(item.get("doc_count", ""))
                row[3].text = str(item.get("count", ""))
                row[4].text = "、".join(item.get("documents") or [])
        else:
            doc.add_paragraph("暂无跨文档共现实体。")

        doc.add_heading("填写准确率", level=2)
        doc.add_paragraph("系统支持模板审核后确认写入；准确率以已匹配字段 / 模板字段数计算。")

        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def _build_pdf_report(self) -> bytes:
        documents = DocumentDAO.get_all()
        entities = EntityDAO.get_all()
        cross_rows = EntityDAO.get_cross_document_entities(min_documents=2, limit=100)
        type_counts = Counter(entity.entity_type for entity in entities)
        lines = [
            "DocFusion Full Workflow Report",
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            "",
            f"Documents: {len(documents)}",
            f"Entities: {len(entities)}",
            "",
            "Entity Type Counts:",
        ]
        lines.extend(f"- {entity_type}: {count}" for entity_type, count in sorted(type_counts.items()))
        lines.append("")
        lines.append("Cross Document Entities:")
        if cross_rows:
            for item in cross_rows:
                docs = ", ".join(item.get("documents") or [])
                lines.append(
                    f"- {item.get('type', '')}: {item.get('value', '')} "
                    f"({item.get('doc_count', '')} docs, {item.get('count', '')} mentions) {docs}"
                )
        else:
            lines.append("- None")
        return self._simple_pdf(lines)

    @staticmethod
    def _simple_pdf(lines: list[str]) -> bytes:
        def esc(value: str) -> str:
            return str(value).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

        text_ops = ["BT", "/F1 11 Tf", "50 790 Td"]
        for index, line in enumerate(lines[:48]):
            if index:
                text_ops.append("0 -16 Td")
            text_ops.append(f"({esc(line)}) Tj")
        text_ops.append("ET")
        stream = "\n".join(text_ops).encode("latin-1", errors="replace")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        ]
        pdf = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for number, obj in enumerate(objects, start=1):
            offsets.append(len(pdf))
            pdf.extend(f"{number} 0 obj\n".encode("ascii"))
            pdf.extend(obj)
            pdf.extend(b"\nendobj\n")
        xref = len(pdf)
        pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
        pdf.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        pdf.extend(
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
        )
        return bytes(pdf)
