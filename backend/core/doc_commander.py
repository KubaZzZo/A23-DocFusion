"""Natural-language document command parsing and execution.

Execution strategies (in priority order):
1. Rule-based parser for simple deterministic formatting commands (fast, no LLM/codex)
2. Codex CLI as primary path for complex commands (writes+executes python-docx code)
3. LLM JSON parsing as fallback when Codex CLI unavailable
"""
import asyncio
import json
import logging
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from datetime import datetime
from typing import Literal
from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator
from llm import get_llm
from llm.base import strip_json_code_fence
from llm.prompt_safety import UNTRUSTED_INPUT_NOTICE, wrap_untrusted_input
from config import DATA_DIR

logger = logging.getLogger(__name__)

BACKUP_DIR = DATA_DIR / "backups"
BACKUP_DIR.mkdir(exist_ok=True)

CODEX_DOC_COMMAND_PROMPT = """You are a document automation agent. The user wants to modify a .docx file using natural language.

Read the document from {doc_path}, apply the user's requested changes using python-docx, and save it back to the same path.

Rules:
- Use `from docx import Document; from docx.shared import Pt, Cm, RGBColor, Inches; from docx.enum.text import WD_ALIGN_PARAGRAPH`
- Read the document first to understand its structure (paragraph count, tables, styles)
- Apply changes precisely as requested
- For paragraph indexing: paragraph 0 = first paragraph, paragraph 1 = second, etc.
- For font names use exact names like "宋体", "黑体", "SimSun", "SimHei", "Arial"
- For colors use RGBColor(r, g, b) with values 0-255
- For font sizes use Pt(size)
- When setting font_name on a run, also set run.font.name and for Chinese fonts set run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', font_name) to handle East Asian characters
- Handle table rows separately from paragraphs
- Save the document after changes
- Print a brief summary of what was done

Common color mappings:
- 红色/red → RGBColor(220, 38, 38)
- 蓝色/blue → RGBColor(37, 99, 235)
- 绿色/green → RGBColor(22, 163, 74)
- 黑色/black → RGBColor(0, 0, 0)
- 灰色/gray → RGBColor(107, 114, 128)

Common alignment: left=WD_ALIGN_PARAGRAPH.LEFT, center=WD_ALIGN_PARAGRAPH.CENTER, right=WD_ALIGN_PARAGRAPH.RIGHT, justify=WD_ALIGN_PARAGRAPH.JUSTIFY

User request:
{user_input}

Document info:
{doc_info}"""

COMMAND_PARSE_PROMPT = """You parse natural-language document editing requests into strict JSON commands.
Supported actions:
1. format - formatting changes: bold, italic, underline, font_size, font_name, color, alignment.
2. edit - content edits: insert, delete, replace.
3. extract - extract text, tables, or headings.
4. find_replace - find and replace text.
5. structure - add headings or paragraphs.

Return only JSON in this shape:
{
  "action": "format|edit|extract|find_replace|structure",
  "target": "paragraph|table_row|text|tables|headings|all",
  "params": {},
  "description": "short operation description"
}

Examples:
Make the second paragraph bold -> {"action":"format","target":"paragraph","params":{"index":1,"bold":true},"description":"make paragraph 2 bold"}
Replace every occurrence of company with enterprise -> {"action":"find_replace","target":"all","params":{"find":"company","replace":"enterprise"},"description":"replace text globally"}
Extract all tables -> {"action":"extract","target":"tables","params":{},"description":"extract all tables"}
"""


class _StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid")


class FormatParams(_StrictModel):
    target: Literal["paragraph", "table_row"] | None = None
    index: int = Field(default=0, ge=0)
    bold: bool | None = None
    italic: bool | None = None
    underline: bool | None = None
    font_size: float | None = Field(default=None, gt=0)
    font_name: str | None = None
    color: tuple[int, int, int] | None = None
    alignment: Literal["left", "center", "right", "justify"] | None = None

    @model_validator(mode="after")
    def _validate_color(self):
        if self.color is not None and any(value < 0 or value > 255 for value in self.color):
            raise ValueError("color must be three integers between 0 and 255")
        return self


class EditParams(_StrictModel):
    operation: Literal["replace", "insert", "delete"] = "replace"
    index: int | None = Field(default=None, ge=0)
    text: str = ""

    @model_validator(mode="after")
    def _validate_required_fields(self):
        if self.operation in {"replace", "delete"} and self.index is None:
            raise ValueError("index must be a non-negative integer")
        return self


class FindReplaceParams(_StrictModel):
    find: str = Field(min_length=1)
    replace: str = ""


class ExtractParams(_StrictModel):
    target: Literal["text", "tables", "headings"] | None = None


class StructureParams(_StrictModel):
    operation: Literal["add_heading", "add_paragraph"] = "add_paragraph"
    text: str = ""
    level: int = Field(default=1, ge=1, le=9)


class CommandSchema(_StrictModel):
    action: Literal["format", "edit", "find_replace", "extract", "structure"]
    target: str | None = None
    params: dict = Field(default_factory=dict)
    description: str = ""


PARAM_SCHEMAS = {
    "format": FormatParams,
    "edit": EditParams,
    "find_replace": FindReplaceParams,
    "extract": ExtractParams,
    "structure": StructureParams,
}


class DocCommander:
    """Document command executor with Codex CLI as primary execution path."""

    def __init__(self, provider: str = None, use_codex: bool = None, llm_client=None):
        self.llm = llm_client or (get_llm(provider) if not use_codex else None)
        self.use_codex = use_codex if use_codex is not None else (
            os.getenv("DOCFUSION_USE_CODEX", "1") == "1"
        )

    # ── primary public API ────────────────────────────────────────

    async def execute_command(
        self, user_input: str, doc_path: str, doc_info: str = ""
    ) -> dict:
        """Execute a natural-language document command.

        Tries rule-based path first, then Codex CLI, then LLM fallback.
        """
        from core.document_parser import DocumentParser

        suffix = Path(doc_path).suffix.lower()
        if suffix not in DocumentParser.SUPPORTED_TYPES:
            return {"success": False, "message": f"Unsupported file format: {suffix}"}

        # Write operations require .docx (python-docx library limitation)
        if suffix != ".docx":
            # Only allow read-only extract operations for non-.docx files
            parsed = await self.parse_command(user_input, doc_info)
            if "error" in parsed:
                return {"success": False, "message": parsed["error"]}
            if parsed.get("action") != "extract":
                return {"success": False, "message": f"Only extract operations are supported for {suffix} files. Write operations require .docx format."}
            return self.execute(doc_path, parsed)

        rule_result = self._try_rule_based_execute(user_input, doc_path)
        if rule_result:
            return rule_result

        if self.use_codex:
            try:
                return await self.execute_via_codex(doc_path, user_input, doc_info)
            except Exception as exc:
                logger.warning("Codex execution failed, falling back to LLM: %s", exc)

        parsed = await self.parse_command(user_input, doc_info)
        if "error" in parsed:
            return {"success": False, "message": parsed["error"]}
        return self.execute(doc_path, parsed)

    # ── codex-powered execution (primary complex-command path) ────

    async def execute_via_codex(
        self, doc_path: str, user_input: str, doc_info: str = ""
    ) -> dict:
        """Execute a document command by having Codex CLI write python-docx code."""
        prompt = CODEX_DOC_COMMAND_PROMPT.format(
            doc_path=doc_path,
            user_input=user_input,
            doc_info=doc_info or f"File: {Path(doc_path).name}",
        )
        backup_path = self._backup(doc_path)
        try:
            result = await self._run_codex_with_prompt(prompt, doc_path)
            if result["success"]:
                return result
        except Exception as exc:
            logger.warning("Codex prompt execution failed: %s", exc)
        if backup_path and Path(backup_path).exists():
            shutil.copyfile(backup_path, doc_path)
        return {"success": False, "message": "Codex CLI execution failed"}

    async def _run_codex_with_prompt(self, prompt: str, doc_path: str) -> dict:
        """Execute a prompt via Codex CLI and collect the result."""
        custom_command = os.getenv("DOCFUSION_CODEX_COMMAND", "").strip()
        if custom_command:
            return await asyncio.to_thread(
                self._run_custom_codex, custom_command, prompt, doc_path
            )
        return await asyncio.to_thread(self._run_codex_cli, prompt, doc_path)

    @staticmethod
    def _run_codex_cli(prompt: str, doc_path: str) -> dict:
        """Run the standard codex CLI."""
        try:
            completed = subprocess.run(
                [
                    "codex", "exec",
                    "--ephemeral",
                    "--skip-git-repo-check",
                    "--ignore-rules",
                    "-s", "workspace-write",
                    "-C", str(Path(doc_path).parent),
                    "-",
                ],
                input=prompt,
                text=True,
                capture_output=True,
                encoding="utf-8",
                errors="replace",
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired):
            return {"success": False, "message": "Codex CLI unavailable or timed out"}

        if completed.returncode != 0:
            return {
                "success": False,
                "message": f"Codex CLI exited with code {completed.returncode}",
                "stderr": (completed.stderr or "")[-500:],
            }
        return {
            "success": True,
            "message": "Document modified via Codex CLI",
            "stdout": (completed.stdout or "")[-500:],
        }

    @staticmethod
    def _run_custom_codex(custom_command: str, prompt: str, doc_path: str) -> dict:
        """Run a user-specified codex-compatible command."""
        try:
            completed = subprocess.run(
                [custom_command],
                input=prompt,
                text=True,
                capture_output=True,
                encoding="utf-8",
                errors="replace",
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired):
            return {"success": False, "message": "Custom codex command unavailable or timed out"}
        if completed.returncode != 0:
            return {
                "success": False,
                "message": f"Command exited with code {completed.returncode}",
                "stderr": (completed.stderr or "")[-500:],
            }
        return {
            "success": True,
            "message": "Document modified via custom codex command",
            "stdout": (completed.stdout or "")[-500:],
        }

    # ── rule-based fast path ──────────────────────────────────────

    def _try_rule_based_execute(self, user_input: str, doc_path: str) -> dict | None:
        """Try to handle the full command with deterministic rules.
        Returns None if rules can't handle it, otherwise the result dict.
        """
        parsed = self._parse_rule_based_command(user_input)
        if not parsed:
            return None
        backup_path = self._backup(doc_path)
        try:
            params = dict(parsed.get("params", {}))
            if "target" in parsed and "target" not in params:
                params["target"] = parsed["target"]
            result = self._handle_format(doc_path, params)
            if backup_path:
                result["backup_path"] = str(backup_path)
            return result
        except Exception as e:
            if backup_path and Path(backup_path).exists():
                shutil.copyfile(backup_path, doc_path)
            return None

    # ── legacy parse+execute API (LLM-parsed JSON, kept as fallback) ──

    async def parse_command(self, user_input: str, doc_info: str = "") -> dict:
        """Parse natural-language input into a command JSON object."""
        rule_based = self._parse_rule_based_command(user_input)
        if rule_based:
            return rule_based

        messages = [
            {"role": "system", "content": COMMAND_PARSE_PROMPT},
            {
                "role": "user",
                "content": (
                    f"{UNTRUSTED_INPUT_NOTICE}\n\n"
                    f"document_info:\n{wrap_untrusted_input(doc_info)}\n\n"
                    f"user_command:\n{wrap_untrusted_input(user_input)}"
                ),
            },
        ]
        try:
            result = await self.llm.chat(messages)
        except Exception as exc:
            codex_parsed = await self._parse_with_codex_cli(user_input, doc_info)
            if codex_parsed:
                return self._normalize_command(user_input, codex_parsed)
            raise exc
        try:
            cleaned = strip_json_code_fence(result)
            parsed = json.loads(cleaned)
            return self._normalize_command(user_input, parsed)
        except json.JSONDecodeError:
            codex_parsed = await self._parse_with_codex_cli(user_input, doc_info)
            if codex_parsed:
                return self._normalize_command(user_input, codex_parsed)
            return {"error": "指令解析失败", "raw": result}

    @classmethod
    def _parse_rule_based_command(cls, user_input: str) -> dict | None:
        """Parse common deterministic formatting requests without an LLM."""
        text = (user_input or "").strip()
        if not text:
            return None

        params = cls._read_format_params(text)
        if not params:
            return None

        row_index = cls._read_ordinal_index(text)
        params["index"] = row_index if row_index is not None else 0
        return {
            "action": "format",
            "target": "paragraph",
            "params": params,
            "description": "rule-based paragraph formatting",
        }

    @staticmethod
    def _read_format_params(text: str) -> dict:
        params: dict[str, object] = {}
        lowered = text.lower()
        if "加粗" in text or "粗体" in text or "bold" in lowered:
            params["bold"] = True
        if "取消加粗" in text or "不加粗" in text:
            params["bold"] = False
        if "斜体" in text or "italic" in lowered:
            params["italic"] = True
        if "下划线" in text or "underline" in lowered:
            params["underline"] = True

        size_match = re.search(r"(?:字号|字体大小|font size|size)\D{0,8}(\d{1,2}(?:\.\d+)?)", text, flags=re.IGNORECASE)
        if not size_match:
            size_match = re.search(r"(\d{1,2}(?:\.\d+)?)\s*(?:号|pt|磅)", text, flags=re.IGNORECASE)
        if size_match:
            params["font_size"] = float(size_match.group(1))
        elif "字号加大" in text or "放大字号" in text or "字体加大" in text or "bigger" in lowered:
            params["font_size"] = 18.0

        color = DocCommander._read_color(text)
        if color is not None:
            params["color"] = color

        if "居中" in text or "center" in lowered:
            params["alignment"] = "center"
        elif "右对齐" in text or "right align" in lowered:
            params["alignment"] = "right"
        elif "两端对齐" in text or "justify" in lowered:
            params["alignment"] = "justify"
        elif "左对齐" in text or "left align" in lowered:
            params["alignment"] = "left"
        return params

    @staticmethod
    def _read_ordinal_index(text: str) -> int | None:
        match = re.search(r"第\s*(\d+)\s*(?:行|段|段落|paragraph|line)", text, flags=re.IGNORECASE)
        if match:
            return max(int(match.group(1)) - 1, 0)

        cn_digits = {
            "一": 1,
            "二": 2,
            "两": 2,
            "三": 3,
            "四": 4,
            "五": 5,
            "六": 6,
            "七": 7,
            "八": 8,
            "九": 9,
            "十": 10,
        }
        match = re.search(r"第\s*([一二两三四五六七八九十])\s*(?:行|段|段落)", text)
        if match:
            return cn_digits[match.group(1)] - 1

        lowered = text.lower()
        english_ordinals = {
            "first": 0,
            "second": 1,
            "third": 2,
            "fourth": 3,
            "fifth": 4,
            "sixth": 5,
            "seventh": 6,
            "eighth": 7,
            "ninth": 8,
            "tenth": 9,
        }
        for word, index in english_ordinals.items():
            if re.search(rf"\b{word}\s+(paragraph|line)\b", lowered):
                return index
        return None

    @staticmethod
    def _read_color(text: str) -> tuple[int, int, int] | None:
        lowered = text.lower()
        color_map = {
            "红色": (220, 38, 38),
            "红": (220, 38, 38),
            "red": (220, 38, 38),
            "蓝色": (37, 99, 235),
            "蓝": (37, 99, 235),
            "blue": (37, 99, 235),
            "绿色": (22, 163, 74),
            "绿": (22, 163, 74),
            "green": (22, 163, 74),
            "黑色": (0, 0, 0),
            "黑": (0, 0, 0),
            "black": (0, 0, 0),
            "灰色": (107, 114, 128),
            "灰": (107, 114, 128),
            "gray": (107, 114, 128),
            "grey": (107, 114, 128),
        }
        for name, rgb in color_map.items():
            if name in lowered or name in text:
                return rgb
        if "变色" in text or "改色" in text or "字体颜色" in text or "文字颜色" in text:
            return color_map["红色"]
        return None

    async def _parse_with_codex_cli(self, user_input: str, doc_info: str = "") -> dict | None:
        """Use local Codex CLI as a slow fallback parser for complex commands."""
        try:
            return await asyncio.to_thread(self._run_codex_cli_parser, user_input, doc_info)
        except Exception:
            return None

    @staticmethod
    def _run_codex_cli_parser(user_input: str, doc_info: str = "") -> dict | None:
        prompt = f"""{COMMAND_PARSE_PROMPT}

Return only one JSON object. Do not explain. Do not edit files.

document_info:
{doc_info}

        user_command:
{user_input}
"""
        custom_command = os.getenv("DOCFUSION_CODEX_COMMAND", "").strip()
        if custom_command:
            try:
                completed = subprocess.run(
                    [custom_command],
                    input=prompt,
                    text=True,
                    capture_output=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=90,
                    check=False,
                )
            except (OSError, subprocess.TimeoutExpired):
                return None
            if completed.returncode != 0:
                return None
            raw = (completed.stdout or "").strip()
            return DocCommander._parse_codex_json(raw)

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "codex-command.json"
            try:
                completed = subprocess.run(
                    [
                        "codex",
                        "exec",
                        "--ephemeral",
                        "--skip-git-repo-check",
                        "--ignore-rules",
                        "-s",
                        "read-only",
                        "-C",
                        str(DATA_DIR.parent),
                        "-o",
                        str(output_path),
                        "-",
                    ],
                    input=prompt,
                    text=True,
                    capture_output=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=90,
                    check=False,
                )
            except (OSError, subprocess.TimeoutExpired):
                return None
            if completed.returncode != 0 or not output_path.exists():
                return None
            raw = output_path.read_text(encoding="utf-8", errors="replace").strip()
            return DocCommander._parse_codex_json(raw)

    @staticmethod
    def _parse_codex_json(raw: str) -> dict | None:
        try:
            return json.loads(strip_json_code_fence(raw))
        except json.JSONDecodeError:
            match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
            if not match:
                return None
            try:
                return json.loads(match.group(0))
            except json.JSONDecodeError:
                return None

    @staticmethod
    def _normalize_command(user_input: str, parsed: dict) -> dict:
        """Normalize common LLM command variants."""
        if not isinstance(parsed, dict):
            return parsed

        action = parsed.get("action")
        target = parsed.get("target")
        params = parsed.get("params", {})
        description = parsed.get("description", "")
        text = f"{user_input} {description}"

        if (
            action == "format"
            and target == "paragraph"
            and isinstance(params, dict)
            and "index" in params
            and ("行" in text or "row" in text.lower() or "table" in text.lower())
        ):
            parsed["target"] = "table_row"

        return parsed

    def execute(self, doc_path: str, command: dict) -> dict:
        """Execute a validated document command and back up mutable operations."""
        from core.document_parser import DocumentParser

        suffix = Path(doc_path).suffix.lower()
        if suffix not in DocumentParser.SUPPORTED_TYPES:
            return {"success": False, "message": f"Unsupported file format: {suffix}"}

        # Write operations require .docx (python-docx library limitation)
        action = command.get("action")
        if suffix != ".docx" and action != "extract":
            return {"success": False, "message": f"Only extract operations are supported for {suffix} files. Write operations require .docx format."}

        validation_error = self._validate_command(command)
        if validation_error:
            return {"success": False, "message": validation_error}

        handlers = {
            "format": self._handle_format,
            "edit": self._handle_edit,
            "find_replace": self._handle_find_replace,
            "extract": self._handle_extract,
            "structure": self._handle_structure,
        }
        handler = handlers.get(action)
        if not handler:
            return {"success": False, "message": f"Unsupported action: {action}"}

        backup_path = None if action == "extract" else self._backup(doc_path)

        try:
            params = dict(command.get("params", {}))
            if action in {"format", "extract"} and "target" in command and "target" not in params:
                params["target"] = command["target"]
            result = handler(doc_path, params)
            if backup_path:
                result["backup_path"] = str(backup_path)
            return result
        except Exception as e:
            # Restore the original file when a mutable operation fails.
            if backup_path and Path(backup_path).exists():
                shutil.copyfile(backup_path, doc_path)
            return {"success": False, "message": str(e)}

    @classmethod
    def _validate_command(cls, command: dict) -> str:
        """Validate LLM command output before any file mutation or backup."""
        if not isinstance(command, dict):
            return "command must be an object"

        try:
            parsed = CommandSchema.model_validate(command)
            params = dict(parsed.params)
            if parsed.action in {"format", "extract"} and parsed.target is not None and "target" not in params:
                params["target"] = parsed.target
            PARAM_SCHEMAS[parsed.action].model_validate(params)
        except ValidationError as e:
            return cls._format_validation_error(e)
        except ValueError as e:
            return str(e)
        return ""

    @staticmethod
    def _format_validation_error(error: ValidationError) -> str:
        first = error.errors()[0]
        loc = ".".join(str(part) for part in first.get("loc", ()))
        message = first.get("msg", "invalid command")
        if first.get("type") == "extra_forbidden" and loc:
            return f"{loc} is not allowed"
        return f"{loc}: {message}" if loc else message

    @staticmethod
    def _backup(doc_path: str) -> Path:
        """Create a backup file and return its path."""
        src = Path(doc_path)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = BACKUP_DIR / f"{src.stem}_{timestamp}{src.suffix}"
        shutil.copyfile(doc_path, backup_path)
        return backup_path

    def _handle_format(self, doc_path: str, params: dict) -> dict:
        doc = DocxDocument(doc_path)
        target = params.get("target", "paragraph")

        if target == "table_row":
            row_idx = params.get("index", 0)
            tables = doc.tables
            if not tables:
                return {"success": False, "message": "No table is available for table-row formatting"}
            first_table = tables[0]
            if row_idx >= len(first_table.rows):
                return {"success": False, "message": f"Table row index {row_idx} is out of range"}

            for cell in first_table.rows[row_idx].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if "bold" in params:
                            run.bold = params["bold"]
                        if "italic" in params:
                            run.italic = params["italic"]
                        if "underline" in params:
                            run.underline = params["underline"]
                        if "font_size" in params:
                            run.font.size = Pt(params["font_size"])
                        if "font_name" in params:
                            run.font.name = params["font_name"]
                        if "color" in params:
                            r, g, b = params["color"]
                            run.font.color.rgb = RGBColor(r, g, b)
                if "alignment" in params:
                    align_map = {
                        "left": WD_ALIGN_PARAGRAPH.LEFT,
                        "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT,
                        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                    }
                    for para in cell.paragraphs:
                        para.alignment = align_map.get(params["alignment"], WD_ALIGN_PARAGRAPH.LEFT)

            doc.save(doc_path)
            return {"success": True, "message": f"Formatted table row {row_idx + 1}"}

        idx = params.get("index", 0)
        if idx >= len(doc.paragraphs):
            return {"success": False, "message": f"Paragraph index {idx} is out of range"}

        para = doc.paragraphs[idx]
        for run in para.runs:
            if "bold" in params:
                run.bold = params["bold"]
            if "italic" in params:
                run.italic = params["italic"]
            if "underline" in params:
                run.underline = params["underline"]
            if "font_size" in params:
                run.font.size = Pt(params["font_size"])
            if "font_name" in params:
                run.font.name = params["font_name"]
            if "color" in params:
                r, g, b = params["color"]
                run.font.color.rgb = RGBColor(r, g, b)

        if "alignment" in params:
            align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
            para.alignment = align_map.get(params["alignment"], WD_ALIGN_PARAGRAPH.LEFT)

        doc.save(doc_path)
        return {"success": True, "message": "Formatting completed"}

    def _handle_edit(self, doc_path: str, params: dict) -> dict:
        doc = DocxDocument(doc_path)
        op = params.get("operation", "replace")

        if op == "replace" and "index" in params:
            idx = params["index"]
            if idx < len(doc.paragraphs):
                doc.paragraphs[idx].text = params.get("text", "")
        elif op == "insert":
            doc.add_paragraph(params.get("text", ""))
        elif op == "delete" and "index" in params:
            idx = params["index"]
            if idx < len(doc.paragraphs):
                p = doc.paragraphs[idx]._element
                p.getparent().remove(p)

        doc.save(doc_path)
        return {"success": True, "message": "Edit completed"}

    def _handle_find_replace(self, doc_path: str, params: dict) -> dict:
        doc = DocxDocument(doc_path)
        find_text = params.get("find", "")
        replace_text = params.get("replace", "")
        count = 0
        if not find_text:
            return {"success": False, "message": "Find text cannot be empty"}

        for para in doc.paragraphs:
            if find_text not in para.text:
                continue
            count += self._replace_in_paragraph_runs(para, find_text, replace_text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if find_text not in para.text:
                            continue
                        count += self._replace_in_paragraph_runs(para, find_text, replace_text)

        doc.save(doc_path)
        return {"success": True, "message": f"Replacement completed, {count} occurrence(s) replaced"}

    @staticmethod
    def _replace_in_paragraph_runs(para, find_text: str, replace_text: str) -> int:
        """Replace text at run level while preserving unaffected run formatting."""
        matches = []
        full_text_parts = []
        char_map = []

        for run_idx, run in enumerate(para.runs):
            for char_idx, ch in enumerate(run.text):
                full_text_parts.append(ch)
                char_map.append((run_idx, char_idx))

        full_text = "".join(full_text_parts)
        start = 0
        while True:
            index = full_text.find(find_text, start)
            if index == -1:
                break
            matches.append((index, index + len(find_text)))
            start = index + len(find_text)

        for start_idx, end_idx in reversed(matches):
            start_run_idx, start_char_idx = char_map[start_idx]
            end_run_idx, end_char_idx = char_map[end_idx - 1]
            start_run = para.runs[start_run_idx]

            if start_run_idx == end_run_idx:
                text = start_run.text
                start_run.text = text[:start_char_idx] + replace_text + text[end_char_idx + 1:]
                continue

            start_run.text = start_run.text[:start_char_idx] + replace_text
            for run_idx in range(start_run_idx + 1, end_run_idx):
                para.runs[run_idx].text = ""
            end_run = para.runs[end_run_idx]
            end_run.text = end_run.text[end_char_idx + 1:]

        return len(matches)

    def _handle_extract(self, doc_path: str, params: dict) -> dict:
        suffix = Path(doc_path).suffix.lower()
        if suffix != ".docx":
            from core.document_parser import DocumentParser
            result = DocumentParser.parse(doc_path)
            return {"success": True, "data": result["text"]}

        doc = DocxDocument(doc_path)
        target = params.get("target", "text")

        if target == "tables":
            tables = []
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    rows.append([cell.text for cell in row.cells])
                tables.append(rows)
            return {"success": True, "data": tables}
        elif target == "headings":
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            return {"success": True, "data": headings}
        else:
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            return {"success": True, "data": text}

    def _handle_structure(self, doc_path: str, params: dict) -> dict:
        doc = DocxDocument(doc_path)
        op = params.get("operation", "add_paragraph")

        if op == "add_heading":
            doc.add_heading(params.get("text", ""), level=params.get("level", 1))
        elif op == "add_paragraph":
            doc.add_paragraph(params.get("text", ""))

        doc.save(doc_path)
        return {"success": True, "message": "Structure operation completed"}
