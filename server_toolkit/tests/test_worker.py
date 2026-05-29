import json

from server_toolkit.worker import execute_plan


def test_execute_plan_writes_structured_logs_and_outputs(tmp_path):
    workspace = tmp_path
    (workspace / "input").mkdir()
    (workspace / "work").mkdir()
    (workspace / "output").mkdir()
    (workspace / "logs").mkdir()
    (workspace / "input" / "source.txt").write_text("hello", encoding="utf-8")

    task = {
        "level": "L1",
        "requires_agent": False,
        "inputs": ["input/source.txt"],
        "outputs": ["output/result.txt"],
        "timeout_seconds": 60,
        "steps": [
            {
                "id": "step_1",
                "tool": "docfusion",
                "command": "copy",
                "args": {
                    "input": "input/source.txt",
                    "output": "output/result.txt",
                },
            }
        ],
    }
    task_file = workspace / "task.json"
    task_file.write_text(json.dumps(task), encoding="utf-8")

    result = execute_plan(task_file, workspace)

    assert result.success is True
    assert (workspace / "output" / "result.txt").read_text(encoding="utf-8") == "hello"
    events = [
        json.loads(line)
        for line in (workspace / "logs" / "steps.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert [event["event"] for event in events] == [
        "start",
        "step_started",
        "step_completed",
        "complete",
    ]


def test_execute_plan_blocks_non_docfusion_tools(tmp_path):
    workspace = tmp_path
    (workspace / "input").mkdir()
    (workspace / "work").mkdir()
    (workspace / "output").mkdir()
    (workspace / "logs").mkdir()
    task_file = workspace / "task.json"
    task_file.write_text(
        json.dumps(
            {
                "level": "L1",
                "requires_agent": False,
                "inputs": [],
                "outputs": [],
                "timeout_seconds": 60,
                "steps": [
                    {
                        "id": "bad",
                        "tool": "shell",
                        "command": "echo",
                        "args": {},
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    result = execute_plan(task_file, workspace)

    assert result.success is False
    assert "Unsupported tool" in result.error


def test_execute_plan_runs_fill_template_step(tmp_path):
    from openpyxl import Workbook, load_workbook

    workspace = tmp_path
    (workspace / "input").mkdir()
    (workspace / "work").mkdir()
    (workspace / "output").mkdir()
    (workspace / "logs").mkdir()

    template = workspace / "input" / "template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "name"
    wb.save(template)
    (workspace / "work" / "data.json").write_text(json.dumps({"name": "DocFusion"}), encoding="utf-8")

    task_file = workspace / "task.json"
    task_file.write_text(
        json.dumps(
            {
                "level": "L1",
                "requires_agent": False,
                "inputs": ["input/template.xlsx"],
                "outputs": ["output/filled.xlsx"],
                "timeout_seconds": 60,
                "steps": [
                    {
                        "id": "fill",
                        "tool": "docfusion",
                        "command": "fill-template",
                        "args": {
                            "input": "input/template.xlsx",
                            "data": "work/data.json",
                            "output": "output/filled.xlsx",
                        },
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    result = execute_plan(task_file, workspace)

    assert result.success is True
    filled = load_workbook(workspace / "output" / "filled.xlsx")
    assert filled.active["A2"].value == "DocFusion"
    filled.close()


def test_execute_plan_fails_when_declared_output_is_missing(tmp_path):
    workspace = tmp_path
    (workspace / "input").mkdir()
    (workspace / "work").mkdir()
    (workspace / "output").mkdir()
    (workspace / "logs").mkdir()
    (workspace / "input" / "source.txt").write_text("hello", encoding="utf-8")

    task_file = workspace / "task.json"
    task_file.write_text(
        json.dumps(
            {
                "level": "L1",
                "requires_agent": False,
                "inputs": ["input/source.txt"],
                "outputs": ["output/expected.txt"],
                "timeout_seconds": 60,
                "steps": [
                    {
                        "id": "copy",
                        "tool": "docfusion",
                        "command": "copy",
                        "args": {"input": "input/source.txt", "output": "work/wrong.txt"},
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    result = execute_plan(task_file, workspace)

    assert result.success is False
    assert "declared output missing" in result.error


def test_execute_plan_logs_agent_output(monkeypatch, tmp_path):
    import server_toolkit.worker as worker

    workspace = tmp_path
    (workspace / "input").mkdir()
    (workspace / "work").mkdir()
    (workspace / "output").mkdir()
    (workspace / "logs").mkdir()

    def fake_agent(workspace_path, instruction, on_output=None):
        assert instruction == "generate report"
        if on_output:
            on_output("stdout", "planning document")
            on_output("stdout", "writing output")
        output = workspace / "output" / "generated.docx"
        output.write_bytes(b"PK")
        return {"success": True, "outputs": [str(output)], "stdout": "done"}

    monkeypatch.setattr(worker, "run_codex_agent", fake_agent)
    task_file = workspace / "task.json"
    task_file.write_text(
        json.dumps(
            {
                "level": "L3",
                "requires_agent": True,
                "inputs": [],
                "outputs": ["output/generated.docx"],
                "instruction": "generate report",
                "timeout_seconds": 600,
                "steps": [
                    {
                        "id": "agent",
                        "tool": "docfusion",
                        "command": "agent-generate",
                        "args": {"instruction": "generate report", "inputs": [], "output": "output/generated.docx"},
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    result = execute_plan(task_file, workspace)

    assert result.success is True
    events = [json.loads(line) for line in (workspace / "logs" / "steps.jsonl").read_text(encoding="utf-8").splitlines()]
    agent_lines = [event["line"] for event in events if event["event"] == "agent_output"]
    assert agent_lines == ["planning document", "writing output"]
