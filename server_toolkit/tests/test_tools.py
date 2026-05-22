import json

from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

from server_toolkit.tools import (
    analyze_table,
    convert_file,
    copy_file,
    extract_text,
    fill_template,
    format_docx,
    generate_document,
    split_pdf,
    validate_file,
)


def test_copy_file_preserves_content(tmp_path):
    source = tmp_path / "input.txt"
    target = tmp_path / "output.txt"
    source.write_text("hello", encoding="utf-8")

    copy_file({"input": str(source), "output": str(target)})

    assert target.read_text(encoding="utf-8") == "hello"


def test_extract_text_from_txt(tmp_path):
    source = tmp_path / "input.txt"
    target = tmp_path / "entities.json"
    source.write_text("alpha beta", encoding="utf-8")

    extract_text({"input": str(source), "output": str(target), "type": "text"})

    data = json.loads(target.read_text(encoding="utf-8"))
    assert data["text"] == "alpha beta"
    assert data["file_type"] == "txt"


def test_extract_text_with_schema_extracts_labeled_fields(tmp_path):
    source = tmp_path / "contract.txt"
    target = tmp_path / "entities.json"
    source.write_text("party_a: Acme\namount: 1000\n", encoding="utf-8")

    extract_text(
        {
            "input": str(source),
            "output": str(target),
            "type": "text",
            "schema": "party_a,amount,date",
        }
    )

    data = json.loads(target.read_text(encoding="utf-8"))
    assert data["entities"] == {"party_a": "Acme", "amount": "1000", "date": None}


def test_convert_txt_to_markdown_without_pandoc(tmp_path):
    source = tmp_path / "source.txt"
    output = tmp_path / "source.md"
    source.write_text("Line one\nLine two", encoding="utf-8")

    convert_file({"input": str(source), "to": "md", "output": str(output)})

    assert output.read_text(encoding="utf-8") == "Line one\nLine two"


def test_convert_docx_to_txt_without_pandoc(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "source.txt"
    doc = Document()
    doc.add_paragraph("Hello")
    doc.add_paragraph("World")
    doc.save(source)

    convert_file({"input": str(source), "to": "txt", "output": str(output)})

    assert output.read_text(encoding="utf-8") == "Hello\nWorld"


def test_format_docx_updates_heading_font_size(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "formatted.docx"
    doc = Document()
    doc.add_heading("Title", level=1)
    doc.add_paragraph("Body")
    doc.save(source)

    format_docx(
        {
            "input": str(source),
            "output": str(output),
            "heading_font": "Arial",
            "heading_size": 18,
        }
    )

    formatted = Document(str(output))
    first_run = formatted.paragraphs[0].runs[0]
    assert first_run.font.name == "Arial"
    assert first_run.font.size.pt == 18


def test_format_docx_updates_body_font_line_spacing_and_margins(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "formatted.docx"
    doc = Document()
    doc.add_paragraph("Body")
    doc.save(source)

    format_docx(
        {
            "input": str(source),
            "output": str(output),
            "body_font": "Arial",
            "body_size": 12,
            "line_spacing": 1.5,
            "margin_top_cm": 2.54,
            "margin_left_cm": 3.17,
        }
    )

    formatted = Document(str(output))
    run = formatted.paragraphs[0].runs[0]
    assert run.font.name == "Arial"
    assert run.font.size.pt == 12
    assert formatted.paragraphs[0].paragraph_format.line_spacing == 1.5
    assert round(formatted.sections[0].top_margin.cm, 2) == 2.54
    assert round(formatted.sections[0].left_margin.cm, 2) == 3.17


def test_fill_template_xlsx_writes_matching_fields(tmp_path):
    template = tmp_path / "template.xlsx"
    data = tmp_path / "data.json"
    output = tmp_path / "filled.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "company"
    ws["B1"] = "amount"
    wb.save(template)
    data.write_text(json.dumps({"company": "Acme", "amount": "1000"}), encoding="utf-8")

    fill_template({"input": str(template), "data": str(data), "output": str(output)})

    filled = load_workbook(output)
    ws = filled.active
    assert ws["A2"].value == "Acme"
    assert ws["B2"].value == "1000"
    filled.close()


def test_generate_document_creates_docx_from_json_sections(tmp_path):
    data = tmp_path / "report.json"
    output = tmp_path / "report.docx"
    data.write_text(
        json.dumps(
            {
                "title": "Project Report",
                "sections": [
                    {"heading": "Summary", "body": "This is the summary."},
                    {"heading": "Risk", "body": "Low."},
                ],
            }
        ),
        encoding="utf-8",
    )

    generate_document({"data": str(data), "output": str(output)})

    doc = Document(str(output))
    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert "Project Report" in texts
    assert "Summary" in texts
    assert "This is the summary." in texts


def test_generate_document_creates_txt_from_json_sections(tmp_path):
    data = tmp_path / "report.json"
    output = tmp_path / "report.txt"
    data.write_text(
        json.dumps({"title": "Report", "sections": [{"heading": "A", "body": "B"}]}),
        encoding="utf-8",
    )

    generate_document({"data": str(data), "output": str(output)})

    assert output.read_text(encoding="utf-8") == "Report\n\nA\nB\n"


def test_split_pdf_extracts_page_range(tmp_path):
    source = tmp_path / "source.pdf"
    output = tmp_path / "page2.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    with source.open("wb") as handle:
        writer.write(handle)

    split_pdf({"input": str(source), "output": str(output), "pages": "2"})

    reader = PdfReader(str(output))
    assert len(reader.pages) == 1


def test_merge_docx_appends_documents(tmp_path):
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    output = tmp_path / "merged.docx"
    doc = Document()
    doc.add_paragraph("First")
    doc.save(first)
    doc = Document()
    doc.add_paragraph("Second")
    doc.save(second)

    from server_toolkit.tools import merge_files

    merge_files({"inputs": [str(first), str(second)], "output": str(output)})

    merged = Document(str(output))
    assert [paragraph.text for paragraph in merged.paragraphs if paragraph.text] == ["First", "Second"]


def test_analyze_table_outputs_summary_for_xlsx(tmp_path):
    source = tmp_path / "data.xlsx"
    output = tmp_path / "summary.json"
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "amount"])
    ws.append(["A", 10])
    ws.append(["B", 20])
    wb.save(source)

    analyze_table({"input": str(source), "output": str(output)})

    summary = json.loads(output.read_text(encoding="utf-8"))
    assert summary["sheets"][0]["rows"] == 2
    assert summary["sheets"][0]["columns"] == ["name", "amount"]
    assert summary["sheets"][0]["numeric"]["amount"]["sum"] == 30


def test_validate_file_rejects_disallowed_suffix(tmp_path):
    source = tmp_path / "payload.exe"
    source.write_bytes(b"MZ")

    try:
        validate_file({"input": str(source), "allowed_types": ["txt"]})
    except Exception as exc:
        assert "unsupported file type" in str(exc)
    else:
        raise AssertionError("expected unsupported file type error")


def test_validate_file_rejects_too_large_file(tmp_path):
    source = tmp_path / "large.txt"
    source.write_text("hello", encoding="utf-8")

    try:
        validate_file({"input": str(source), "max_size": 2})
    except Exception as exc:
        assert "file too large" in str(exc)
    else:
        raise AssertionError("expected file too large error")
