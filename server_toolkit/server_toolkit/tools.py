"""Deterministic document tools used by the server worker."""

from __future__ import annotations

import csv
import json
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Any, Callable


class ToolError(RuntimeError):
    """Raised when a deterministic tool cannot complete."""


ToolFunc = Callable[[dict[str, Any]], dict[str, Any]]


def copy_file(args: dict[str, Any]) -> dict[str, Any]:
    source = Path(_required(args, "input"))
    output = Path(_required(args, "output"))
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, output)
    return {"output": str(output)}


def convert_file(args: dict[str, Any]) -> dict[str, Any]:
    source = Path(_required(args, "input"))
    output = Path(_required(args, "output"))
    target = str(args.get("to") or output.suffix.lstrip(".")).lower()
    output.parent.mkdir(parents=True, exist_ok=True)

    if source.suffix.lower() in {".txt", ".md", ".csv"} and target in {"txt", "md", "markdown", "csv"}:
        shutil.copyfile(source, output)
        return {"output": str(output)}
    if source.suffix.lower() == ".docx" and target in {"txt", "md", "markdown"}:
        text = _extract_docx(source)
        output.write_text(text, encoding="utf-8")
        return {"output": str(output), "chars": len(text)}
    if source.suffix.lower() in {".txt", ".md"} and target == "docx":
        return _text_to_docx(source, output)
    if source.suffix.lower() == ".xlsx" and target == "csv":
        return _xlsx_to_csv(source, output)
    if target in {"pdf", "docx", "xlsx", "pptx"}:
        return _libreoffice_convert(source, output, target)
    if target in {"md", "markdown", "html", "txt"}:
        return _pandoc_convert(source, output)
    raise ToolError(f"unsupported conversion target: {target}")


def format_docx(args: dict[str, Any]) -> dict[str, Any]:
    try:
        from docx import Document
        from docx.shared import Cm
        from docx.shared import RGBColor
        from docx.shared import Pt
    except ImportError as exc:
        raise ToolError("python-docx is required for format-docx") from exc

    source = Path(_required(args, "input"))
    output = Path(_required(args, "output"))
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(source))

    heading_font = args.get("heading_font")
    heading_size = args.get("heading_size")
    body_font = args.get("body_font")
    body_size = args.get("body_size")
    line_spacing = args.get("line_spacing")
    first_line_bold = bool(args.get("first_line_bold"))
    first_line_underline = bool(args.get("first_line_underline"))
    first_line_color = args.get("first_line_color")

    margin_map = {
        "margin_top_cm": "top_margin",
        "margin_bottom_cm": "bottom_margin",
        "margin_left_cm": "left_margin",
        "margin_right_cm": "right_margin",
    }
    for section in doc.sections:
        for arg_name, attr_name in margin_map.items():
            if args.get(arg_name) is not None:
                setattr(section, attr_name, Cm(float(args[arg_name])))

    for paragraph in doc.paragraphs:
        style_name = (paragraph.style.name or "").lower()
        is_heading = style_name.startswith("heading")
        if line_spacing is not None and not is_heading:
            paragraph.paragraph_format.line_spacing = float(line_spacing)
        for run in paragraph.runs:
            if is_heading:
                if heading_font:
                    run.font.name = str(heading_font)
                if heading_size:
                    run.font.size = Pt(float(heading_size))
            else:
                if body_font:
                    run.font.name = str(body_font)
                if body_size:
                    run.font.size = Pt(float(body_size))

    if first_line_bold or first_line_underline or first_line_color:
        first_paragraph = next((paragraph for paragraph in doc.paragraphs if paragraph.text.strip()), None)
        if first_paragraph is not None:
            if not first_paragraph.runs:
                first_paragraph.add_run(first_paragraph.text)
            for run in first_paragraph.runs:
                if first_line_bold:
                    run.font.bold = True
                if first_line_underline:
                    run.font.underline = True
                color = _parse_hex_color(first_line_color)
                if color:
                    run.font.color.rgb = RGBColor(*color)

    doc.save(str(output))
    return {"output": str(output)}


def _parse_hex_color(value: Any) -> tuple[int, int, int] | None:
    if not value:
        return None
    text = str(value).strip().lstrip("#")
    if len(text) != 6:
        return None
    try:
        return int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16)
    except ValueError:
        return None


def merge_files(args: dict[str, Any]) -> dict[str, Any]:
    inputs = [Path(item) for item in args.get("inputs", [])]
    output = Path(_required(args, "output"))
    if not inputs:
        raise ToolError("inputs is required")
    output.parent.mkdir(parents=True, exist_ok=True)
    suffix = output.suffix.lower()
    if suffix == ".pdf":
        return _merge_pdf(inputs, output)
    if suffix == ".docx":
        return _merge_docx(inputs, output)
    if suffix == ".txt":
        output.write_text(
            "\n".join(path.read_text(encoding="utf-8") for path in inputs),
            encoding="utf-8",
        )
        return {"output": str(output)}
    raise ToolError(f"unsupported merge output type: {suffix}")


def split_pdf(args: dict[str, Any]) -> dict[str, Any]:
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise ToolError("pypdf is required for PDF split") from exc

    source = Path(_required(args, "input"))
    output = Path(_required(args, "output"))
    output.parent.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(str(source))
    writer = PdfWriter()
    for page_number in _parse_pages(str(_required(args, "pages")), len(reader.pages)):
        writer.add_page(reader.pages[page_number - 1])
    with output.open("wb") as handle:
        writer.write(handle)
    return {"output": str(output)}


def extract_text(args: dict[str, Any]) -> dict[str, Any]:
    source = Path(_required(args, "input"))
    output = Path(_required(args, "output"))
    output.parent.mkdir(parents=True, exist_ok=True)
    suffix = source.suffix.lower()

    if suffix in {".txt", ".md", ".csv"}:
        text = source.read_text(encoding="utf-8")
    elif suffix == ".docx":
        text = _extract_docx(source)
    elif suffix == ".xlsx":
        text = _extract_xlsx(source)
    elif suffix == ".pdf":
        text = _extract_pdf(source)
    else:
        raise ToolError(f"unsupported extract input type: {suffix}")

    payload = {"text": text, "file_type": suffix.lstrip(".")}
    if args.get("schema"):
        payload["entities"] = _extract_schema_entities(text, str(args["schema"]))
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"output": str(output), "chars": len(text)}


def fill_template(args: dict[str, Any]) -> dict[str, Any]:
    source = Path(_required(args, "input"))
    data_path = Path(_required(args, "data"))
    output = Path(_required(args, "output"))
    output.parent.mkdir(parents=True, exist_ok=True)
    data = json.loads(data_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ToolError("data JSON must be an object")
    suffix = source.suffix.lower()
    if suffix == ".xlsx":
        return _fill_xlsx_template(source, data, output)
    if suffix == ".docx":
        return _fill_docx_template(source, data, output)
    raise ToolError(f"unsupported template type: {suffix}")


def generate_document(args: dict[str, Any]) -> dict[str, Any]:
    data_path = Path(_required(args, "data"))
    output = Path(_required(args, "output"))
    output.parent.mkdir(parents=True, exist_ok=True)
    data = json.loads(data_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ToolError("data JSON must be an object")
    suffix = output.suffix.lower()
    if suffix == ".docx":
        return _generate_docx(data, output)
    if suffix in {".txt", ".md"}:
        text = _render_plain_document(data)
        output.write_text(text, encoding="utf-8")
        return {"output": str(output), "chars": len(text)}
    raise ToolError(f"unsupported generate output type: {suffix}")


def analyze_table(args: dict[str, Any]) -> dict[str, Any]:
    source = Path(_required(args, "input"))
    output = Path(_required(args, "output"))
    output.parent.mkdir(parents=True, exist_ok=True)
    suffix = source.suffix.lower()
    if suffix == ".xlsx":
        payload = _analyze_xlsx(source)
    elif suffix == ".csv":
        payload = _analyze_csv(source)
    else:
        raise ToolError(f"unsupported analyze input type: {suffix}")
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"output": str(output)}


def ocr_image(args: dict[str, Any]) -> dict[str, Any]:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise ToolError("pytesseract and Pillow are required for ocr") from exc

    source = Path(_required(args, "input"))
    output = Path(_required(args, "output"))
    lang = str(args.get("lang") or "chi_sim+eng")
    output.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as image:
        text = pytesseract.image_to_string(image, lang=lang)
    output.write_text(text, encoding="utf-8")
    return {"output": str(output), "chars": len(text)}


def validate_file(args: dict[str, Any]) -> dict[str, Any]:
    source = Path(_required(args, "input"))
    if not source.exists():
        raise ToolError(f"file not found: {source}")
    max_size = args.get("max_size")
    size = source.stat().st_size
    if max_size is not None and size > int(max_size):
        raise ToolError(f"file too large: {size} bytes")
    suffix = source.suffix.lower().lstrip(".")
    allowed_types = args.get("allowed_types")
    if isinstance(allowed_types, str):
        allowed_types = [item.strip() for item in allowed_types.split(",") if item.strip()]
    if allowed_types and suffix not in {str(item).lower().lstrip(".") for item in allowed_types}:
        raise ToolError(f"unsupported file type: {suffix}")
    signature = _detect_signature(source)
    return {"path": str(source), "size": size, "suffix": f".{suffix}", "signature": signature}


TOOL_REGISTRY: dict[str, ToolFunc] = {
    "copy": copy_file,
    "convert": convert_file,
    "format-docx": format_docx,
    "fill-template": fill_template,
    "merge": merge_files,
    "split": split_pdf,
    "extract": extract_text,
    "generate": generate_document,
    "analyze": analyze_table,
    "ocr": ocr_image,
    "validate": validate_file,
}


def run_tool(command: str, args: dict[str, Any]) -> dict[str, Any]:
    func = TOOL_REGISTRY.get(command)
    if not func:
        raise ToolError(f"Unsupported docfusion command: {command}")
    return func(args)


def _required(args: dict[str, Any], key: str) -> str:
    value = args.get(key)
    if not value:
        raise ToolError(f"{key} is required")
    return str(value)


def _xlsx_to_csv(source: Path, output: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ToolError("openpyxl is required for xlsx to csv conversion") from exc

    wb = load_workbook(str(source), read_only=True, data_only=True)
    ws = wb.active
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if cell is None else cell for cell in row])
    wb.close()
    return {"output": str(output)}


def _text_to_docx(source: Path, output: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ToolError("python-docx is required for text to docx conversion") from exc

    doc = Document()
    for line in source.read_text(encoding="utf-8").splitlines():
        if line.strip():
            doc.add_paragraph(line.strip())
    doc.save(str(output))
    return {"output": str(output)}


def _libreoffice_convert(source: Path, output: Path, target: str) -> dict[str, Any]:
    exe = shutil.which("libreoffice") or shutil.which("soffice")
    if not exe:
        raise ToolError("LibreOffice is required for this conversion")
    subprocess.run(
        [exe, "--headless", "--convert-to", target, "--outdir", str(output.parent), str(source)],
        check=True,
        capture_output=True,
        text=True,
    )
    generated = output.parent / f"{source.stem}.{target}"
    if generated != output and generated.exists():
        generated.replace(output)
    if not output.exists():
        raise ToolError(f"conversion did not produce output: {output}")
    return {"output": str(output)}


def _pandoc_convert(source: Path, output: Path) -> dict[str, Any]:
    exe = shutil.which("pandoc")
    if not exe:
        raise ToolError("Pandoc is required for this conversion")
    subprocess.run([exe, str(source), "-o", str(output)], check=True, capture_output=True, text=True)
    return {"output": str(output)}


def _merge_pdf(inputs: list[Path], output: Path) -> dict[str, Any]:
    try:
        from pypdf import PdfWriter
    except ImportError as exc:
        raise ToolError("pypdf is required for PDF merge") from exc

    writer = PdfWriter()
    for path in inputs:
        writer.append(str(path))
    with output.open("wb") as handle:
        writer.write(handle)
    return {"output": str(output)}


def _merge_docx(inputs: list[Path], output: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ToolError("python-docx is required for docx merge") from exc

    merged = Document(str(inputs[0]))
    for source in inputs[1:]:
        doc = Document(str(source))
        for paragraph in doc.paragraphs:
            if paragraph.text:
                merged.add_paragraph(paragraph.text)
        for table in doc.tables:
            if not table.rows:
                continue
            new_table = merged.add_table(rows=len(table.rows), cols=len(table.columns))
            for row_index, row in enumerate(table.rows):
                for col_index, cell in enumerate(row.cells):
                    new_table.cell(row_index, col_index).text = cell.text
    merged.save(str(output))
    return {"output": str(output)}


def _parse_pages(value: str, page_count: int) -> list[int]:
    pages: list[int] = []
    for part in value.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start_text, end_text = part.split("-", 1)
            start = int(start_text)
            end = int(end_text)
            pages.extend(range(start, end + 1))
        else:
            pages.append(int(part))
    if not pages:
        raise ToolError("pages must select at least one page")
    for page in pages:
        if page < 1 or page > page_count:
            raise ToolError(f"page out of range: {page}")
    return pages


def _fill_xlsx_template(source: Path, data: dict[str, Any], output: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ToolError("openpyxl is required for xlsx template fill") from exc

    wb = load_workbook(str(source))
    filled = 0
    for ws in wb.worksheets:
        headers = [cell.value for cell in ws[1]]
        for col_index, header in enumerate(headers, start=1):
            key = str(header).strip() if header is not None else ""
            if key in data:
                ws.cell(row=2, column=col_index, value=_safe_cell_value(data[key]))
                filled += 1
    wb.save(str(output))
    wb.close()
    return {"output": str(output), "filled": filled}


def _fill_docx_template(source: Path, data: dict[str, Any], output: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ToolError("python-docx is required for docx template fill") from exc

    doc = Document(str(source))
    filled = 0
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            marker = "{{" + str(key) + "}}"
            if marker in paragraph.text:
                paragraph.text = paragraph.text.replace(marker, str(value))
                filled += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for key, value in data.items():
                    marker = "{{" + str(key) + "}}"
                    if marker in cell.text:
                        cell.text = cell.text.replace(marker, str(value))
                        filled += 1
    doc.save(str(output))
    return {"output": str(output), "filled": filled}


def _generate_docx(data: dict[str, Any], output: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ToolError("python-docx is required for docx generation") from exc

    doc = Document()
    title = str(data.get("title") or "").strip()
    if title:
        doc.add_heading(title, level=0)
    for section in _read_sections(data):
        heading = str(section.get("heading") or "").strip()
        body = str(section.get("body") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        if body:
            for paragraph in body.splitlines():
                if paragraph.strip():
                    doc.add_paragraph(paragraph.strip())
    doc.save(str(output))
    return {"output": str(output)}


def _render_plain_document(data: dict[str, Any]) -> str:
    parts: list[str] = []
    title = str(data.get("title") or "").strip()
    if title:
        parts.append(title)
    section_blocks = []
    for section in _read_sections(data):
        heading = str(section.get("heading") or "").strip()
        body = str(section.get("body") or "").strip()
        block_parts = []
        if heading:
            block_parts.append(heading)
        if body:
            block_parts.append(body)
        if block_parts:
            section_blocks.append("\n".join(block_parts))
    parts.extend(section_blocks)
    return "\n\n".join(parts) + "\n"


def _read_sections(data: dict[str, Any]) -> list[dict[str, Any]]:
    sections = data.get("sections") or []
    if not isinstance(sections, list):
        raise ToolError("sections must be a list")
    result = []
    for item in sections:
        if not isinstance(item, dict):
            raise ToolError("each section must be an object")
        result.append(item)
    return result


def _analyze_xlsx(source: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ToolError("openpyxl is required for xlsx analysis") from exc

    wb = load_workbook(str(source), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        sheets.append(_analyze_rows(ws.title, rows))
    wb.close()
    return {"sheets": sheets}


def _analyze_csv(source: Path) -> dict[str, Any]:
    with source.open("r", encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))
    return {"sheets": [_analyze_rows(source.stem, rows)]}


def _analyze_rows(name: str, rows: list[Any]) -> dict[str, Any]:
    if not rows:
        return {"name": name, "rows": 0, "columns": [], "numeric": {}}
    headers = [str(cell) if cell is not None else "" for cell in rows[0]]
    data_rows = rows[1:]
    numeric: dict[str, dict[str, float | int]] = {}
    for index, header in enumerate(headers):
        values = []
        for row in data_rows:
            try:
                value = row[index]
            except IndexError:
                continue
            if isinstance(value, (int, float)):
                values.append(float(value))
            elif isinstance(value, str):
                try:
                    values.append(float(value))
                except ValueError:
                    pass
        if values:
            numeric[header] = {
                "count": len(values),
                "sum": sum(values),
                "min": min(values),
                "max": max(values),
            }
    return {"name": name, "rows": len(data_rows), "columns": headers, "numeric": numeric}


def _extract_schema_entities(text: str, schema: str) -> dict[str, str | None]:
    fields = [item.strip() for item in schema.split(",") if item.strip()]
    entities: dict[str, str | None] = {field: None for field in fields}
    lines = text.splitlines()
    for field in fields:
        prefixes = (field + ":", field + "：")
        for line in lines:
            stripped = line.strip()
            for prefix in prefixes:
                if stripped.lower().startswith(prefix.lower()):
                    entities[field] = stripped[len(prefix):].strip() or None
                    break
            if entities[field] is not None:
                break
    return entities


def _safe_cell_value(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _extract_docx(source: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ToolError("python-docx is required for docx extraction") from exc

    doc = Document(str(source))
    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(source: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ToolError("openpyxl is required for xlsx extraction") from exc

    wb = load_workbook(str(source), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[Sheet {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _extract_pdf(source: Path) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise ToolError("PyMuPDF is required for PDF extraction") from exc

    parts = []
    with fitz.open(str(source)) as doc:
        for page in doc:
            text = page.get_text()
            if text.strip():
                parts.append(text)
    return "\n".join(parts)


def _detect_signature(source: Path) -> str:
    with source.open("rb") as handle:
        head = handle.read(8)
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        try:
            with zipfile.ZipFile(source) as archive:
                names = set(archive.namelist())
            if "word/document.xml" in names:
                return "docx"
            if "xl/workbook.xml" in names:
                return "xlsx"
            if "ppt/presentation.xml" in names:
                return "pptx"
        except zipfile.BadZipFile:
            return "zip"
        return "zip"
    if head.startswith(b"\xff\xd8"):
        return "jpg"
    if head.startswith(b"\x89PNG"):
        return "png"
    return "unknown"
